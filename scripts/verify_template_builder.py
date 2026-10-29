"""Verify actual company Word template plus SOP inside the Builder."""
from pathlib import Path
from tempfile import TemporaryDirectory
import shutil,sys,time
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from PySide6.QtWidgets import QApplication
from PySide6.QtTest import QTest
from sop_capture.storage import ProjectStore
from sop_capture.templates import TemplateLibrary
from sop_capture.ui.main_window import MainWindow
from sop_capture.ui.theme import stylesheet

def main():
    app=QApplication([]);out=Path('.superpowers/template-builder');out.mkdir(parents=True,exist_ok=True)
    with TemporaryDirectory() as temporary:
        root=Path(temporary);shutil.copytree('demo/Permit Management',root/'project');store=ProjectStore(root/'project');project=store.load()
        source=root/'company.docx';doc=Document();doc.sections[0].header.paragraphs[0].text='EXAMPLE COMPANY | CONTROLLED DOCUMENT'
        doc.add_heading('{{title}}',0)
        table=doc.add_table(rows=3,cols=2);table.style='Table Grid'
        for row,(label,value) in zip(table.rows,[('Document owner','Not yet assigned'),('Reference','{{reference}}'),('Version','{{version}}')]):row.cells[0].text=label;row.cells[1].text=value
        doc.add_heading('Company requirements',1);doc.add_paragraph('Complete the permit and obtain supervisor approval before starting work.')
        doc.add_heading('Standard operating procedure',1);doc.add_paragraph('{{steps}}')
        doc.sections[0].footer.paragraphs[0].text='Company controlled copy | {{reference}}';doc.save(source)
        library=TemplateLibrary(root/'library');profile=library.save('Company SOP layout',{},source_docx=source)
        project.template=library.snapshot(profile['id'],store.root);project.company_information='';window=MainWindow(template_library=library);window.set_project(project,store)
        editor=window.template_text_editor;index=next(i for i,row in enumerate(editor.records) if row['text']=='Not yet assigned');editor.picker.setCurrentIndex(index);editor.text.setPlainText('Site operations team')
        window.resize(1440,900);window.builder_modes.setCurrentIndex(1);window.show();window.show_page(2);window.builder_edit_tabs.setCurrentIndex(2)
        deadline=time.monotonic()+50
        while time.monotonic()<deadline:
            app.processEvents();time.sleep(.02)
            if window.builder_preview.document and not window.builder_preview.running:break
        if not window.builder_preview.document:raise RuntimeError(window.builder_preview.status.text())
        for dark in (False,True):
            window.setStyleSheet(stylesheet(dark));app.processEvents();QTest.qWait(1200);window.grab().save(str(out/f'loaded-template-{ "dark" if dark else "light"}.png'))
        preview=window.builder_preview
        (out/'combined-document.pdf').write_bytes(bytes(preview.source.data()))
        print('Rendered pages:',preview.document.pageCount())
        preview.page.setValue(min(2,preview.document.pageCount()));app.processEvents();QTest.qWait(1200);window.grab().save(str(out/'sop-in-template.png'))
        window.close()
    print(out.resolve())

if __name__=='__main__':main()
