"""Read company Word styles and create an empty reusable page layout."""
from collections import Counter
from io import BytesIO
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def normal_name(name):
    return re.sub(r'[^a-z0-9]', '', name.casefold())


def numbered(style):
    seen=set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        identifiers=style.element.xpath('./w:pPr/w:numPr/w:numId/@w:val')
        if identifiers:return identifiers[0]!='0'
        style=style.base_style
    return False


def unnumber(paragraph):
    properties=paragraph._p.get_or_add_pPr()
    existing=properties.find(qn('w:numPr'))
    if existing is not None:properties.remove(existing)
    num=OxmlElement('w:numPr');identifier=OxmlElement('w:numId');identifier.set(qn('w:val'),'0');num.append(identifier);properties.append(num)


def heading_level(style):
    seen=set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        levels=style.element.xpath('./w:pPr/w:outlineLvl/@w:val')
        if levels:return int(levels[0])+1 if int(levels[0])<9 else None
        style=style.base_style
    return None


def learn_styles(document):
    styles=[s for s in document.styles if s.type==WD_STYLE_TYPE.PARAGRAPH]
    names={s.name:s for s in styles}
    used=Counter(p.style.name for p in document.paragraphs if p.text.strip())
    def choose(candidates,fallback):
        return max(candidates,key=lambda s:(s.name in used,used[s.name],s.name!='Normal')).name if candidates else fallback
    def outline_level(style):
        seen=set()
        while style is not None and style.style_id not in seen:
            seen.add(style.style_id)
            levels=style.element.xpath('./w:pPr/w:outlineLvl/@w:val')
            if levels:return int(levels[0])
            style=style.base_style
        return 9
    ordinary=[s for s in styles if s.name in used and outline_level(s)==9 and not
              any(term in normal_name(s.name) for term in ('heading','title','caption','image','bullet','number','header','footer','list'))]
    fallback=choose(ordinary,'Normal' if 'Normal' in names else styles[0].name)
    bodies=[s for s in styles if any(term in normal_name(s.name) for term in ('bodytext','textbody','normaltext'))]
    bodies=[s for s in bodies if s.name in used or not s.builtin]
    body=choose(bodies,fallback)
    prefix=normal_name(body)
    for suffix in ('bodytext','textbody','normaltext'):
        if prefix.endswith(suffix):prefix=prefix[:-len(suffix)];break
    def heading(level,fallback):
        candidates=[s for s in styles if outline_level(s)==level-1 or re.search(r'(heading|head|titre|uberschrift)'+str(level)+r'$',normal_name(s.name))]
        company=[s for s in candidates if prefix and normal_name(s.name).startswith(prefix)]
        return choose(company or candidates,fallback)
    title=heading(1,'Title' if 'Title' in names else body)
    step=heading(2,title)
    captions=[s for s in styles if 'caption' in normal_name(s.name)]
    company_captions=[s for s in captions if prefix and normal_name(s.name).startswith(prefix)]
    return {'body_style':body,'heading_style':step,'title_style':title,
            'caption_style':choose(company_captions or captions,body)}


def prepare_template(path):
    from .company_fields import xp
    document=Document(path)
    settings=learn_styles(document)
    from docx.text.paragraph import Paragraph
    body=document.element.body;remove_level=None;notes=False;last_heading=''
    anchor=next((p._p for p in document.paragraphs if re.fullmatch(r'\s*\{\{\s*steps\s*\}\}\s*',p.text)),None)
    for element in list(body):
        if element is anchor:continue
        text=' '.join(xp(element,'.//w:t/text()')).strip()
        name=normal_name(text)
        paragraph=Paragraph(element,document._body) if element.tag==qn('w:p') else None
        style=normal_name(paragraph.style.name) if paragraph is not None else ''
        match=re.search(r'(?:heading|head|titre)([1-9])$',style)
        level=heading_level(paragraph.style) if paragraph is not None else None
        if level is None and match:level=int(match[1])
        if level is not None:last_heading=name
        guidance=name in {'usingthistemplate','templateinstructions','howtousethistemplate'}
        demonstration=bool(re.fullmatch(r'(?:meh)?heading[1-9]',name))
        if remove_level is not None and level is not None and level<=remove_level and not demonstration:
            remove_level=None
        if notes and paragraph is not None and (level is not None or 'subheading' in style or name in {'tableofcontents','contents'}):notes=False
        if notes and element.tag==qn('w:tbl'):notes=False
        if guidance or demonstration:
            remove_level=level or 1
            if anchor is None:
                anchor=OxmlElement('w:p');run=OxmlElement('w:r');node=OxmlElement('w:t');node.text='{{steps}}';run.append(node);anchor.append(run);element.addprevious(anchor)
        if name.startswith('tousetemplateread') or name in {'tousetemplate','instructionsforcompletingthistemplate'}:notes=True
        discard=guidance or demonstration or remove_level is not None or notes
        if element.tag==qn('w:tbl'):
            values=[normal_name(''.join(xp(cell,'.//w:t/text()'))) for cell in xp(element,'.//w:tc')]
            if values and 'tableheading' in values and all(value in {'tableheading','tabletext','tablebulletlevel1tablebulletlevel2'} for value in values):discard=True
        if discard:
            # Section properties belong to the company layout, even when a
            # guidance paragraph carrying a section break is removed.
            if paragraph is not None and element.xpath('./w:pPr/w:sectPr'):
                for child in list(element):
                    if child.tag!=qn('w:pPr'):element.remove(child)
            elif element.tag!=qn('w:sectPr'):body.remove(element)
        elif 'lorem ipsum' in text.casefold():
            # Keep its surrounding introduction/section; only sample text goes.
            for sample in xp(element,'self::w:p|.//w:p'):
                if 'lorem ipsum' not in ''.join(xp(sample,'.//w:t/text()')).casefold():continue
                nodes=xp(sample,'.//w:t')
                for node in nodes:node.text=''
                target={'introduction':'purpose','purpose':'purpose','scope':'scope','prerequisites':'prerequisites'}.get(last_heading)
                if target and nodes:nodes[0].text='{{'+target+'}}'
    if anchor is None and not any(''.join(xp(p,'.//w:t/text()')).strip()=='{{steps}}' for p in body if p.tag==qn('w:p')):
        paragraph=document.add_paragraph('{{steps}}');anchor=paragraph._p
    update=document.settings.element.find(qn('w:updateFields'))
    if update is None:update=OxmlElement('w:updateFields');document.settings.element.append(update)
    update.set(qn('w:val'),'true')
    # Keep TOC field instructions, but not cached entries referencing deleted
    # demonstration sections. Word rebuilds these when opening/exporting.
    stack=[]
    for node in body.iter():
        if node.tag==qn('w:fldChar'):
            kind=node.get(qn('w:fldCharType'))
            if kind=='begin':stack.append({'code':'','result':False})
            elif kind=='separate' and stack:stack[-1]['result']=True
            elif kind=='end' and stack:stack.pop()
        elif node.tag==qn('w:instrText') and stack:stack[-1]['code']+=node.text or ''
        elif node.tag==qn('w:t') and any(field['result'] and re.match(r'\s*TOC\b',field['code'],re.I) for field in stack):node.text=''
    data=BytesIO();document.save(data)
    return data.getvalue(),settings


def has_structure(data):
    from .company_fields import xp
    document=Document(BytesIO(data))
    return any(''.join(xp(el,'.//w:t/text()')).strip() not in {'','{{steps}}'} for el in document.element.body if el.tag!=qn('w:sectPr'))


def fill_company_fields(document, project):
    """Recognised human-facing title prompts, including split Word runs."""
    from .word_export import stories
    for paragraph in stories(document):
        prompts=re.findall(r'[<\[]\s*Insert\s+Document\s+Title\s*[>\]]',''.join(paragraph._p.xpath('.//w:t/text()')),re.I)
        for prompt in prompts+['{{title}}', '{{sop_title}}']:
            replace_prompt(paragraph,prompt,project.title)
    from .company_fields import fill_form
    fill_form(document,project)


def replace_prompt(paragraph,prompt,value):
    replace_element_text(paragraph._p,prompt,value)


def replace_element_text(element,prompt,value):
    from .company_fields import xp
    nodes=xp(element,'.//w:t')
    text=''.join(node.text or '' for node in nodes)
    for match in reversed(list(re.finditer(re.escape(prompt),text))):
        offset=0;touched=[]
        for node in nodes:
            end=offset+len(node.text or '')
            if offset<match.end() and end>match.start():touched.append((node,offset))
            offset=end
        for i,(node,start) in enumerate(touched):
            current=node.text or '';left=max(0,match.start()-start);right=min(len(current),match.end()-start)
            node.text=current[:left]+(value if i==0 else '')+current[right:]
            node.set(qn('xml:space'),'preserve')
