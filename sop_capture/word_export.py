"""Word template inspection and local document generation without template code."""
from copy import copy
from io import BytesIO
from pathlib import Path
import re
import tempfile
import zipfile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree
from PIL import Image

from .imaging import render_step
from .storage import ProjectStore, atomic_bytes
from .templates import appearance

FIELDS = ("title", "sop_title", "company", "reference", "version", "author", "reviewer", "approver", "status", "created", "updated", "purpose", "scope", "prerequisites", "revision_history", "steps", "logo")
TOKEN = re.compile(r"\{\{\s*([a-zA-Z_][a-zA-Z0-9_]*)\s*\}\}")


def stories(document):
    for part in document.part.package.parts:
        element = getattr(part, "element", None)
        if element is not None and etree.QName(element).localname in {"document", "hdr", "ftr"}:
            for paragraph in element.xpath(".//w:p"):
                yield Paragraph(paragraph, part)


def validate_fields(element):
    unsafe = re.compile(r"\s*(DDEAUTO|DDE|INCLUDETEXT|INCLUDEPICTURE|LINK|DATABASE|IMPORT|RD)\b", re.I)
    def check(instruction):
        if unsafe.match(instruction): raise ValueError("Remove external or active Word fields before importing")
    stack = []
    for node in element.iter():
        local = etree.QName(node).localname
        if local == "fldSimple": check(node.get(qn("w:instr"), ""))
        elif local == "fldChar":
            kind = node.get(qn("w:fldCharType"))
            if kind == "begin": stack.append("")
            elif kind == "separate" and stack: check(stack[-1])
            elif kind == "end" and stack: check(stack.pop())
        elif local == "instrText":
            if stack: stack[-1] += node.text or ""
            else: check(node.text or "")
    for incomplete in stack: check(incomplete)


def inspect_docx(path):
    path = Path(path)
    try:
        if path.suffix.lower() != ".docx" or path.stat().st_size > 30_000_000:
            raise ValueError("Choose a .docx template under 30 MB")
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        with zipfile.ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > 150_000_000:
                raise ValueError("Word template is too large when expanded")
            for name in archive.namelist():
                if "vbaproject" in name.lower(): raise ValueError("Macro documents are not supported")
                if "/embeddings/" in name or "/activeX/" in name: raise ValueError("Templates with embedded active content are not supported")
                if name.endswith(".xml"):
                    data = archive.read(name)
                    if b"<!DOCTYPE" in data: raise ValueError("Document type declarations are not supported")
                    element = etree.fromstring(data, parser)
                    validate_fields(element)
                if name.endswith(".rels"):
                    root = etree.fromstring(archive.read(name), parser)
                    for relationship in root:
                        if relationship.get("TargetMode") == "External" and not relationship.get("Type", "").endswith("/hyperlink"):
                            raise ValueError("Remove externally linked images or attached templates before importing")
            content_types = archive.read("[Content_Types].xml")
            if b"macroEnabled" in content_types: raise ValueError("Macro documents are not supported")
        document = Document(path)
        texts = [p.text for p in stories(document)]
        tokens = sorted({m.group(1) for text in texts for m in TOKEN.finditer(text)})
        unknown = set(tokens)-set(FIELDS)
        if unknown: raise ValueError("Unsupported template fields: " + ", ".join(sorted(unknown)))
        anchors = [p for p in document.paragraphs if re.fullmatch(r"\s*\{\{\s*steps\s*\}\}\s*", p.text)]
        if "steps" in tokens and (len(anchors) != 1 or sum(m.group(1) == "steps" for text in texts for m in TOKEN.finditer(text)) != 1):
            raise ValueError("Use {{steps}} once, alone in a body paragraph outside tables/headers")
        if any(m.group(1) == "logo" and not re.fullmatch(r"\s*\{\{\s*logo\s*\}\}\s*", text) for text in texts for m in TOKEN.finditer(text)):
            raise ValueError("Place {{logo}} alone in its own paragraph")
        return {"styles": sorted(s.name for s in document.styles if s.type == WD_STYLE_TYPE.PARAGRAPH),
                "fields": tokens, "has_steps": bool(anchors)}
    except (OSError, zipfile.BadZipFile, KeyError, etree.XMLSyntaxError) as exc:
        raise ValueError(f"Cannot read Word template: {exc}") from exc


def replace_tokens(paragraph, values):
    # Replace in reverse order to preserve unaffected runs and their formatting,
    # including placeholders that Word split across several runs.
    text = paragraph.text
    for match in reversed(list(TOKEN.finditer(text))):
        key = match.group(1)
        if key not in values: continue
        runs = [Run(element, paragraph) for element in paragraph._p.xpath(".//w:r")]
        offset = 0
        touched = []
        for run in runs:
            end = offset+len(run.text)
            if offset < match.end() and end > match.start(): touched.append((run, offset, end))
            offset = end
        if not touched: continue
        for index, (run, start, end) in enumerate(touched):
            left = max(0, match.start()-start)
            right = min(len(run.text), match.end()-start)
            run.text = run.text[:left]+(str(values[key]) if index == 0 else "")+run.text[right:]


def export_docx(project, root, output):
    if project.document_file:
        source = ProjectStore(root).resolve(project.document_file)
        if source == Path(output).resolve():
            raise ValueError("Choose a separate export file; the saved document is kept unchanged.")
        atomic_bytes(output, source.read_bytes())
        return Path(output)
    project = copy(project)
    settings = appearance(project.template.get("settings", {}))
    store = ProjectStore(root)
    source = project.template.get("docx")
    prepared=bool(source and project.template.get('prepared'))
    if source:
        source_path = store.resolve(source)
        if Path(output).resolve()==source_path.resolve():raise ValueError("Choose a separate export file; the original Word template is kept unchanged.")
        info = inspect_docx(source_path)
        document = Document(source_path)
        from .template_text import apply_edits
        apply_edits(document, source_path, project.template_text_edits)
        if prepared:
            from .word_template import fill_company_fields
            fill_company_fields(document,project)
    else:
        document = Document()
        info = {"has_steps": False}
    styles = {s.name for s in document.styles if s.type == WD_STYLE_TYPE.PARAGRAPH}
    def style(key, fallback):
        name = settings[key]
        if name not in styles:
            raise ValueError(f"Word style '{name}' is missing. Choose a style on the Templates page and apply again.")
        return name
    heading = style("heading_style", "Heading 2")
    body = style("body_style", "Normal")
    caption = style("caption_style", "Caption")
    project.company = project.company or settings["company"]
    values = {key: getattr(project, key) for key in FIELDS if key not in {"steps", "logo"} and hasattr(project, key)}
    values["sop_title"] = project.title
    logo = project.template.get("logo") or project.logo
    if not source:
        font = {"Vera": "Bitstream Vera Sans", "Helvetica": "Arial", "Times-Roman": "Times New Roman", "Courier": "Courier New"}[settings["font"]]
        for name, size, colour in ((body, settings["body_size"], settings["text_color"]), (heading, settings["heading_size"], settings["heading_color"]), ("Title", settings["title_size"], settings["heading_color"])):
            document.styles[name].font.name = font
            document.styles[name].font.size = Pt(size)
            document.styles[name].font.color.rgb = RGBColor.from_string(colour.lstrip("#"))
        orientation = project.orientation if settings["orientation"] == "project" else settings["orientation"]
        for section in document.sections:
            section.page_width, section.page_height = (Mm(297), Mm(210)) if orientation == "landscape" else (Mm(210), Mm(297))
            section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Mm(settings["margin_mm"])
            def label(text):
                for key in ("title", "reference", "version", "company"): text = text.replace("{"+key+"}", getattr(project, key))
                return text
            section.header.paragraphs[0].text = label(settings["header"] or project.reference+" | "+project.status)
            footer = section.footer.paragraphs[0]
            footer.text = label(settings["footer"])+" | Page "
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE")
            footer._p.append(field)
    section_width = min(s.page_width-s.left_margin-s.right_margin for s in document.sections)
    max_height = min(s.page_height-s.top_margin-s.bottom_margin for s in document.sections)*.55
    width_inches = section_width/914400*settings["image_width_percent"]/100
    for paragraph in list(stories(document)):
        if re.fullmatch(r"\s*\{\{\s*logo\s*\}\}\s*", paragraph.text):
            paragraph.clear()
            if logo:
                with Image.open(store.resolve(logo)) as image:
                    ratio = min(1.8/image.width, .8/image.height)
                    paragraph.add_run().add_picture(str(store.resolve(logo)), width=Inches(image.width*ratio), height=Inches(image.height*ratio))
        else:
            replace_tokens(paragraph, values)
    anchor = next((p for p in document.paragraphs if re.fullmatch(r"\s*\{\{\s*steps\s*\}\}\s*", p.text)), None)
    generated = []
    def paragraph(text="", name=body, keep=False):
        p = document.add_paragraph(text, style=name)
        p.paragraph_format.keep_with_next = keep
        generated.append(p._p)
        return p
    structured=prepared and project.template.get('structured',False)
    if not info["has_steps"] or (prepared and not structured):
        if source and not prepared: paragraph().paragraph_format.page_break_before = True
        if project.company: paragraph(project.company, caption)
        if logo and not source:
            with Image.open(store.resolve(logo)) as image:
                scale = min(1.8/image.width, .8/image.height)
                paragraph().add_run().add_picture(str(store.resolve(logo)), width=Inches(image.width*scale), height=Inches(image.height*scale))
        title_style=style('title_style',heading) if prepared else ("Title" if "Title" in styles else heading)
        title_paragraph=paragraph(project.title,title_style)
        if prepared:
            from .word_template import unnumber
            unnumber(title_paragraph)
        paragraph(f"{project.reference} | Version {project.version} | {project.status}\nAuthor: {project.author}")
        if settings["cover"] and not source:
            for label, value in (("Reviewer", project.reviewer), ("Approver", project.approver), ("Created", project.created[:10]), ("Updated", project.updated[:10])):
                paragraph(f"{label}: {value or '—'}")
            paragraph().paragraph_format.page_break_before = True
        for label, value in (("Purpose", project.purpose), ("Scope", project.scope), ("Prerequisites", project.prerequisites)):
            if value:
                paragraph(label, heading, True)
                paragraph(value)
    from .company_information import word_information
    from .company_fields import xp
    preceding=anchor._p.getprevious() if anchor is not None else None
    while preceding is not None and not ''.join(xp(preceding,'.//w:t/text()')).strip():preceding=preceding.getprevious()
    existing_procedure=''.join(xp(preceding,'.//w:t/text()')).strip().casefold() if preceding is not None else ''
    if structured and existing_procedure not in {'procedure','standard operating procedure','sop'}:
        paragraph('Procedure',style('title_style',heading),True)
    word_information(document, project.company_information, body, heading, generated)
    previous_section = ""
    for number, step in enumerate(project.steps, 1):
        section_heading = None
        if step.section and step.section != previous_section:
            section_heading = paragraph(step.section, heading, True)
        previous_section = step.section
        from .word_template import numbered
        step_text=step.title if prepared and numbered(document.styles[heading]) else f"{number:02d}  {step.title}"
        title = paragraph(step_text, heading, True)
        if settings["page_per_step"]:
            (section_heading if section_heading is not None else title).paragraph_format.page_break_before = True
        if step.image:
            try:
                image = render_step(root, step, number, project)
                data = BytesIO()
                image.save(data, "PNG")
                data.seek(0)
                ratio = min(width_inches/image.width, (max_height/914400)/image.height)
                paragraph(keep=True).add_run().add_picture(data, width=Inches(image.width*ratio), height=Inches(image.height*ratio))
                paragraph(f"Step {number} screenshot", caption, True)
            except (OSError, ValueError): paragraph("Image unavailable — replace the screenshot in the editor.")
        paragraph(step.description)
        if step.notes: paragraph("Notes / warnings: "+step.notes)
    if not project.steps: paragraph("No procedural steps have been added.")
    if (not info["has_steps"] or (prepared and not structured)) and project.revision_history:
        paragraph("Revision history", heading, True)
        paragraph(project.revision_history)
    if anchor:
        for element in generated: anchor._p.addprevious(element)
        anchor._p.getparent().remove(anchor._p)
    data = BytesIO()
    document.save(data)
    atomic_bytes(output, data.getvalue())
    return Path(output)


def word_pdf(docx_path, output):
    """Use a separate desktop Word automation instance; never attach to user docs."""
    from .word_session import WordSession
    session = None
    try:
        session = WordSession(docx_path, visible=False, read_only=True)
        session.document.ExportAsFixedFormat(str(Path(output).resolve()), 17)
    except Exception as exc:
        raise RuntimeError("Document PDF export requires a working desktop Microsoft Word installation. You can export DOCX instead. " + str(exc)) from exc
    finally:
        if session is not None: session.close()
    if not Path(output).is_file(): raise RuntimeError("Word did not create a PDF")
    return Path(output)


def create_example(path):
    document = Document()
    for style in document.styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH:
            style.font.name = "Arial"
    for name, size, colour in (("Normal", 10, "183247"), ("Title", 26, "000000"), ("Heading 1", 16, "15344B"), ("Heading 2", 13, "15344B"), ("Caption", 9, "526777")):
        document.styles[name].font.size = Pt(size)
        document.styles[name].font.color.rgb = RGBColor.from_string(colour)
        for border in document.styles[name]._element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    for section in document.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin = section.right_margin = Mm(20)
    document.add_paragraph("{{company}}", style="Subtitle")
    document.add_paragraph("{{title}}", style="Title")
    document.add_paragraph("Reference: {{reference}}    Version: {{version}}\nAuthor: {{author}}    Status: {{status}}")
    document.add_paragraph("{{logo}}")
    document.add_paragraph("Purpose", style="Heading 1")
    document.add_paragraph("{{purpose}}")
    document.add_paragraph("Scope", style="Heading 1")
    document.add_paragraph("{{scope}}")
    document.add_paragraph("Procedure", style="Heading 1")
    document.add_paragraph("{{steps}}")
    document.add_paragraph("Revision history", style="Heading 1")
    document.add_paragraph("{{revision_history}}")
    document.sections[0].header.paragraphs[0].text = "{{reference}} | {{status}}"
    data = BytesIO()
    document.save(data)
    atomic_bytes(path, data.getvalue())
    return Path(path)
