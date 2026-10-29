"""Isolated proof that desktop Word can edit a synthetic document in a Qt pane."""
from pathlib import Path
from tempfile import TemporaryDirectory
import time
import pythoncom
import win32com.client
import win32gui,win32process
from docx import Document
from PySide6.QtWidgets import QApplication,QWidget,QVBoxLayout,QLabel
from PySide6.QtGui import QWindow

app=QApplication([])
existing=set()
def find(hwnd,_):
 if win32gui.GetClassName(hwnd)=='OpusApp':existing.add(win32process.GetWindowThreadProcessId(hwnd)[1])
win32gui.EnumWindows(find,None)
pythoncom.CoInitialize();word=None;doc=None;window=None;foreign=None
with TemporaryDirectory(prefix='sop-word-embed-') as root:
 path=Path(root)/'synthetic.docx';source=Document();source.add_heading('Editable document test',0);source.add_paragraph('Original company information');source.save(path)
 try:
  word=win32com.client.DispatchEx('Word.Application');word.Visible=False;word.DisplayAlerts=0;word.AutomationSecurity=3
  doc=word.Documents.Open(str(path),ReadOnly=False,AddToRecentFiles=False,Visible=True)
  hwnd=int(doc.ActiveWindow.Hwnd);pid=win32process.GetWindowThreadProcessId(hwnd)[1]
  if pid in existing:raise RuntimeError('Word reused an existing process; embedding not attempted.')
  window=QWidget();window.resize(1100,800);box=QVBoxLayout(window);box.addWidget(QLabel('Synthetic Word document — direct editing test'))
  foreign=QWindow.fromWinId(hwnd)
  if foreign is None:raise RuntimeError('Foreign Word window could not be hosted.')
  container=QWidget.createWindowContainer(foreign,window);box.addWidget(container,1);window.show();word.Visible=True
  deadline=time.monotonic()+3
  while time.monotonic()<deadline:app.processEvents();pythoncom.PumpWaitingMessages();time.sleep(.02)
  print('Embedded:',win32gui.GetParent(hwnd)!=0,'Class:',win32gui.GetClassName(hwnd),'PID:',pid,flush=True)
  doc.Range(doc.Content.End-1,doc.Content.End-1).InsertAfter('\nSaved inside the hosted Word editor.')
  doc.Save();print('Editable and saved:', 'Saved inside' in doc.Content.Text,flush=True)
  Path('.superpowers/visual').mkdir(parents=True,exist_ok=True)
  from PIL import ImageGrab
  rect=win32gui.GetWindowRect(int(window.winId()));ImageGrab.grab(bbox=rect).save('.superpowers/visual/word-embedded-prototype.png')
 finally:
  if doc is not None:doc.Close(False)
  if word is not None:word.Quit()
  if window is not None:window.close()
  app.processEvents()
pythoncom.CoUninitialize()
