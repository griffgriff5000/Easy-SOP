from docx import Document
from sop_capture.models import Step
from sop_capture.storage import ProjectStore
from sop_capture.templates import TemplateLibrary
from sop_capture.ui.main_window import MainWindow
from PySide6.QtWidgets import QDialog
from sop_capture.ui.new_project_dialog import NewProjectDialog


def library_template(tmp_path):
    source=tmp_path/'company.docx';document=Document();document.add_paragraph('Company cover');document.add_paragraph('{{steps}}');document.save(source)
    library=TemplateLibrary(tmp_path/'library');profile=library.save('Company procedure',{},source_docx=source)
    return library,profile


def test_new_projects_reuse_library_template_in_separate_folders(tmp_path,monkeypatch):
    library,profile=library_template(tmp_path);parent=tmp_path/'projects';parent.mkdir()
    chosen=[]
    def choose(dialog):
        chosen.append([dialog.template.itemText(i) for i in range(dialog.template.count())])
        dialog.title.setText('Permit process');dialog.location.setText(str(parent));dialog.template.setCurrentIndex(dialog.template.findData(profile['id']))
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(NewProjectDialog,'exec',choose)
    window=MainWindow(template_library=library);window.refresh_windows=lambda:None
    window.settings.fields['logo'].blockSignals(True);window.settings.fields['logo'].setText('')
    window.new_project();first=window.store;first_file=first.resolve(window.project.template['docx'])
    assert window.project.template['id']==profile['id']
    window.project.steps.append(Step(title='First project only'));window.flush_save()
    window.new_project();second=window.store
    assert first.root!=second.root and first.root.parent==second.root.parent==parent
    assert first.root.name=='SOP - Permit process' and second.root.name=='SOP - Permit process (2)'
    assert len(first.load().steps)==1 and not window.project.steps
    second_file=second.resolve(window.project.template['docx'])
    assert first_file!=second_file and first_file.read_bytes()==second_file.read_bytes()
    second_file.write_bytes(b'project copy changed')
    assert first_file.read_bytes()==library.asset(profile,'docx').read_bytes()
    assert library.load(profile['id'])['name']=='Company procedure'
    assert chosen==[['Standard SOP','Company procedure'],['Standard SOP','Company procedure']]
    window.close();window.deleteLater()


def test_cancel_template_selection_does_not_create_or_switch_project(tmp_path,monkeypatch):
    library,profile=library_template(tmp_path);window=MainWindow(template_library=library)
    monkeypatch.setattr(NewProjectDialog,'exec',lambda self:QDialog.DialogCode.Rejected)
    window.new_project()
    assert window.project is None and not (tmp_path/'SOP - Cancelled project').exists()
    window.close();window.deleteLater()
