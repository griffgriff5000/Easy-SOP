"""Render a revision example and both review themes without touching user files."""
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.shared import Inches
from PIL import Image
from PySide6.QtCore import QStandardPaths
from PySide6.QtWidgets import QApplication
from sop_capture.revisions import RevisionWorkspace
from sop_capture.revision_export import build_document
from sop_capture.ui.main_window import MainWindow
from sop_capture.templates import TemplateLibrary
from sop_capture.designs import AnnotationLibrary
from sop_capture.recent_projects import RecentProjects
from sop_capture.ui.theme import stylesheet


def main():
    import uuid
    QStandardPaths.setTestModeEnabled(True);app=QApplication([]);out=Path('.superpowers/revision')/uuid.uuid4().hex[:8];out.mkdir(parents=True)
    source=out/'Rev A.docx';template=out/'New template.docx'
    image=out/'example.png';Image.new('RGB',(400,150),'#087f8c').save(image)
    doc=Document();doc.add_heading('Equipment setup — Rev A',0);doc.add_heading('Responsibilities',1);doc.add_paragraph('The supervisor confirms that the equipment is safe before use.')
    table=doc.add_table(rows=2,cols=2);table.style='Table Grid';table.cell(0,0).text='Role';table.cell(0,1).text='Responsibility';table.cell(1,0).text='Supervisor';table.cell(1,1).text='Inspect equipment'
    doc.add_heading('Method',1);doc.add_paragraph('Check the power supply.',style='List Number');doc.add_paragraph('Select Start.',style='List Number');doc.add_picture(str(image),width=Inches(3));doc.save(source)
    doc=Document();doc.sections[0].header.paragraphs[0].text='MEH | CONTROLLED PROCEDURE';doc.add_heading('{{title}}',0);doc.add_paragraph('Revision {{revision}}');doc.add_heading('Roles and responsibilities',1);doc.add_heading('Procedure',1);doc.sections[0].footer.paragraphs[0].text='Company controlled copy';doc.save(template)
    work=RevisionWorkspace.create(out/'workspace',source,template,'B')
    for e in work.elements:
        if e['kind']=='heading':work.exclude(e['id'],'New template supplies the headings')
        else:work.approve(e['id'],work.destinations[1 if e.get('section')=='Responsibilities' else 2]['id'])
    build_document(work,out/'Rev B.docx')
    window=MainWindow(TemplateLibrary(out/'library'),RecentProjects(out/'recent.json'),AnnotationLibrary(out/'annotations'));window.revision_page.bind(work);window.show_page(7);window.show()
    for dark in (False,True):
        window.setStyleSheet(stylesheet(dark));theme='dark' if dark else 'light'
        for width,height in ((1000,760),(1440,950)):
            window.resize(width,height);window.revision_page.sources.setCurrentRow(3);window.revision_page.right_tabs.setCurrentIndex(0);app.processEvents();window.grab().save(str(out/f'review-{theme}-{width}.png'))
            window.revision_page.right_tabs.setCurrentIndex(2);app.processEvents();window.grab().save(str(out/f'preview-{theme}-{width}.png'))
    window.close()
    from sop_capture.word_export import word_pdf
    import pypdfium2
    word_pdf(out/'Rev B.docx',out/'Rev B.pdf')
    with pypdfium2.PdfDocument(str(out/'Rev B.pdf')) as pdf:
        for i in range(len(pdf)):
            page=pdf[i];bitmap=page.render(scale=1.3);bitmap.to_pil().save(out/f'output-page-{i+1}.png');bitmap.close();page.close()
    print(out.resolve())

if __name__=='__main__':main()
