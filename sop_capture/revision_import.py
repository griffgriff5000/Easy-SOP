"""Loss-aware, entirely local source inspection. Original inputs are read-only."""
from pathlib import Path
from hashlib import sha256
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from .company_fields import xp
from .word_template import heading_level


def validate_word_input(path):
    """Keep the same active-content boundary as SOP template imports."""
    import zipfile
    from .word_export import validate_fields
    if Path(path).stat().st_size>50_000_000:raise ValueError('Choose a Word document under 50 MB.')
    with zipfile.ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist())>200_000_000:raise ValueError('Word document is too large when expanded.')
        parser=etree.XMLParser(resolve_entities=False,no_network=True)
        for name in archive.namelist():
            if 'vbaproject' in name.lower() or '/activeX/' in name or '/embeddings/' in name:raise ValueError('Remove macros or embedded active objects before importing this Word document.')
            if name.endswith(('.xml','.rels')):
                data=archive.read(name)
                if b'<!DOCTYPE' in data:raise ValueError('Document type declarations are not supported.')
                root=etree.fromstring(data,parser);validate_fields(root)
                if name.endswith('.rels'):
                    for rel in root:
                        if rel.get('TargetMode')=='External' and not rel.get('Type','').endswith('/hyperlink'):raise ValueError('Remove externally linked content before importing this Word document.')


def element_text(element):
    def text(paragraph):
        values=[]
        for node in paragraph.iter():
            if node.tag in {qn('w:t'),qn('m:t')}:values.append(node.text or '')
            elif node.tag==qn('w:tab'):values.append('\t')
            elif node.tag in {qn('w:br'),qn('w:cr')}:values.append('\n')
        return ''.join(values)
    return '\n'.join(text(p) for p in xp(element,'self::w:p | .//w:p')).strip()


def inspect_source(path,assets):
    path=Path(path);assets=Path(assets);assets.mkdir(parents=True,exist_ok=True)
    if path.suffix.lower()=='.pdf':return inspect_pdf(path,assets)
    if path.suffix.lower()!='.docx':raise ValueError('Choose a Word .docx or PDF source document.')
    doc=Document(path);items=[];warnings=[];section='Document start'
    for index,node in enumerate(doc.element.body):
        if node.tag==qn('w:sectPr'):continue
        text=element_text(node);kind='text';level=None;rows=None
        if node.tag==qn('w:p'):
            paragraph=Paragraph(node,doc._body);level=heading_level(paragraph.style)
            if level:kind='heading';section=text
        if node.tag==qn('w:tbl'):
            kind='table';rows=[[c.text for c in row.cells] for row in Table(node,doc._body).rows]
        elif xp(node,'.//w:drawing | .//w:pict'):kind='image'
        if not text and not xp(node,'.//w:drawing | .//w:pict | .//w:tbl | .//m:oMath | .//w:object | .//w:footnoteReference | .//w:endnoteReference'):continue
        notices=[]
        for expression,label in [('.//w:object','Embedded object'),('.//w:footnoteReference | .//w:endnoteReference','Footnote/endnote reference'),('.//w:ins | .//w:del','Tracked changes'),('.//w:fldChar | .//w:fldSimple','Word field/cross-reference'),('.//w:sectPr','Source section layout'),('.//w:bookmarkStart','Bookmark'),('.//w:txbxContent','Text box')]:
            if xp(node,expression):notices.append(label+' needs review')
        # Charts and other non-image internal relationships cannot be flattened safely.
        blocked=[]
        for expression,label in [('.//m:oMath | .//m:oMathPara','Equation'),('.//w:footnoteReference | .//w:endnoteReference','Footnote/endnote'),('.//w:commentRangeStart | .//w:commentRangeEnd | .//w:commentReference','Comment'),('.//w:ins | .//w:del','Tracked changes')]:
            if xp(node,expression):blocked.append(label+': transfer manually in Word to preserve dependent content')
        for attr in xp(node,'.//@r:id | .//@r:embed | .//@r:link'):
            rel=doc.part.rels.get(str(attr))
            if rel is None or (not rel.is_external and not rel.reltype.endswith('/image')):blocked.append('Unsupported linked object; transfer manually in Word')
        if node.tag not in {qn('w:p'),qn('w:tbl')}:blocked.append('Grouped Word content control; transfer manually in Word')
        previews=[]
        for rid in xp(node,'.//@r:embed'):
            rel=doc.part.rels.get(str(rid))
            if rel and not rel.is_external and rel.reltype.endswith('/image'):
                image=rel.target_part;filename=sha256(image.blob).hexdigest()[:20]+'.'+image.partname.ext
                (assets/filename).write_bytes(image.blob);previews.append(str(assets/filename))
        item={'id':f'w{index}','kind':kind,'text':text,'section':section,'level':level,'rows':rows,'xml':etree.tostring(node,encoding='unicode'),'images':previews,'warnings':notices+blocked,'blocked':bool(blocked)}
        items.append(item)
    for section_index,sec in enumerate(doc.sections):
        for name,part in [('header',sec.header),('footer',sec.footer)]:
            if any(p.text.strip() for p in part.paragraphs) or part.tables:
                warnings.append(f'Source {name}, section {section_index+1}: retained in original; new template supplies its own. Review separately.')
    if not items:raise ValueError('The source Word document contains no transferable body content.')
    return items,warnings


def inspect_pdf(path,assets):
    import pdfplumber
    import pypdfium2 as pdfium
    items=[];warnings=['PDF reading order, table boundaries and image regions require visual review.']
    with pdfplumber.open(path) as pdf, pdfium.PdfDocument(str(path)) as rendered:
        for number,page in enumerate(pdf.pages,1):
            native=rendered[number-1];bitmap=native.render(scale=1.5);picture=bitmap.to_pil().copy();bitmap.close();native.close()
            image_path=assets/f'page-{number}.png';picture.save(image_path)
            tables=page.find_tables();regions=[]
            for ti,table in enumerate(tables):
                rows=[[str(value or '') for value in row] for row in table.extract()]
                if not rows:continue
                regions.append(table.bbox)
                items.append({'id':f'p{number}t{ti}','kind':'table','rows':rows,'text':'\n'.join(' | '.join(row) for row in rows),'page':number,'section':f'PDF page {number}','top':table.bbox[1],'images':[],'warnings':['Extracted PDF table: check merged cells and column order.']})
            outside=page.filter(lambda obj:not any(b[0]<=obj.get('x0',-1)<b[2] and b[1]<=obj.get('top',-1)<b[3] for b in regions))
            text=outside.extract_text() or ''
            if text.strip():
                items.append({'id':f'p{number}text','kind':'text','text':text,'page':number,'section':f'PDF page {number}','top':0,'images':[],'warnings':['PDF text: verify reading order and line breaks.']})
            for ii,image in enumerate(page.images):
                if not page.chars:continue  # entire scanned page is offered below
                bbox=(max(0,image['x0']),max(0,image['top']),min(page.width,image['x1']),min(page.height,image['bottom']))
                if bbox[2]<=bbox[0] or bbox[3]<=bbox[1]:continue
                crop=picture.crop(tuple(round(v*1.5) for v in bbox));target=assets/f'page-{number}-image-{ii}.png';crop.save(target)
                items.append({'id':f'p{number}image{ii}','kind':'image','text':'PDF image region','page':number,'section':f'PDF page {number}','top':bbox[1],'images':[str(target)],'warnings':['Image is a rendered PDF region; inspect its boundary.']})
            # Always retain a full-page evidence alternative; it is explicitly excluded initially when text exists.
            scanned=not page.chars
            items.append({'id':f'p{number}page','kind':'page','text':f'Page {number} artwork'+(' — scan/no text layer' if scanned else ' — alternative to extracted elements'),'page':number,'section':f'PDF page {number}','top':float('inf'),'images':[str(image_path)],'alternative':not scanned,'warnings':['Scanned/no-text page: transfer artwork or type a transcription. Automatic OCR is not installed.'] if scanned else ['Full-page alternative: avoid approving both this and extracted content.']})
            if scanned:warnings.append(f'Page {number} is a scan or has no text layer; manual transcription or page artwork is required.')
    return sorted(items,key=lambda e:(e.get('page',0),e.get('top',0))),warnings


def destinations(path):
    doc=Document(path);result=[{'id':'start','label':'Document start','index':-1,'kind':'start'}]
    for index,node in enumerate(doc.element.body):
        if node.tag!=qn('w:p'):continue
        paragraph=Paragraph(node,doc._body);text=paragraph.text.strip()
        if text and (heading_level(paragraph.style) or '{{steps}}' in text):
            result.append({'id':f'd{index}','label':text,'index':index,'kind':'section'})
    result.append({'id':'end','label':'End of document','index':len(doc.element.body)-1,'kind':'end'})
    return result
