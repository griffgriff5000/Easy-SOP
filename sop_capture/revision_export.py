"""Build only user-approved transfers on a byte-for-byte template snapshot."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement,parse_xml
from docx.shared import Inches
from .company_fields import xp
from types import SimpleNamespace
from .word_template import learn_styles,fill_company_fields,heading_level,replace_element_text
from .storage import atomic_bytes


def _replace_text(element,text):
    nodes=xp(element,'.//w:t')
    if nodes:
        nodes[0].text=text;nodes[0].set(qn('xml:space'),'preserve')
        for node in nodes[1:]:node.text=''
    else:
        run=OxmlElement('w:r');node=OxmlElement('w:t');node.text=text;run.append(node);element.append(run)


def _numbering(node,source,target,mapping):
    for num in xp(node,'.//w:numPr/w:numId'):
        old=num.get(qn('w:val'))
        if old=='0':continue
        if old not in mapping:
            src=source.part.numbering_part.element;dst=target.part.numbering_part.element
            originals=src.xpath(f'./w:num[@w:numId="{int(old)}"]')
            if not originals:raise ValueError('A source list numbering definition is missing; transfer this element manually.')
            numcopy=deepcopy(originals[0]);abstract=numcopy.find(qn('w:abstractNumId'));aid=abstract.get(qn('w:val'))
            abstracts=src.xpath(f'./w:abstractNum[@w:abstractNumId="{int(aid)}"]')
            if not abstracts:raise ValueError('A source list format is missing.')
            acopy=deepcopy(abstracts[0]);newaid=max([int(n.get(qn('w:abstractNumId'))) for n in dst.xpath('./w:abstractNum')]+[-1])+1
            acopy.set(qn('w:abstractNumId'),str(newaid))
            for linked in xp(acopy,'.//w:pStyle | .//w:styleLink | .//w:numStyleLink'):linked.getparent().remove(linked)
            dst.insert(0,acopy);newnum=max([int(n.get(qn('w:numId'))) for n in dst.xpath('./w:num')]+[0])+1
            numcopy.set(qn('w:numId'),str(newnum));abstract.set(qn('w:val'),str(newaid));dst.append(numcopy);mapping[old]=str(newnum)
        num.set(qn('w:val'),mapping[old])


def _style(node,source,target,settings):
    for paragraph in xp(node,'self::w:p | .//w:p'):
        props=paragraph.find(qn('w:pPr'))
        if props is None:props=OxmlElement('w:pPr');paragraph.insert(0,props)
        existing=props.find(qn('w:pStyle'));old=existing.get(qn('w:val')) if existing is not None else 'Normal'
        level=None
        for style in source.styles:
            if style.style_id==old:level=heading_level(style);break
        name=settings['body_style']
        if level:
            candidates=[s for s in target.styles if s.type==1 and heading_level(s)==level]
            name=candidates[0].name if candidates else settings['title_style'] if level==1 else settings['heading_style']
        if existing is None:existing=OxmlElement('w:pStyle');props.insert(0,existing)
        if level:
            for number in xp(props,'./w:numPr'):props.remove(number)
        if level is None and props.find(qn('w:numPr')) is None:
            original_style=next((s for s in source.styles if s.style_id==old),None);seen=set()
            while original_style is not None and original_style.style_id not in seen:
                seen.add(original_style.style_id);numbering=xp(original_style.element,'./w:pPr/w:numPr')
                if numbering:props.append(deepcopy(numbering[0]));break
                original_style=original_style.base_style
        existing.set(qn('w:val'),target.styles[name].style_id)
        # Keep emphasis, superscripts and table structure, while inheriting company typography.
        for property_node in xp(paragraph,'.//w:rPr/w:rFonts | .//w:rPr/w:sz | .//w:rPr/w:szCs | .//w:rPr/w:color'):
            property_node.getparent().remove(property_node)
        for section in xp(paragraph,'./w:pPr/w:sectPr'):section.getparent().remove(section)
    for table in xp(node,'self::w:tbl | .//w:tbl'):
        props=table.find(qn('w:tblPr'))
        if props is None:props=OxmlElement('w:tblPr');table.insert(0,props)
        old=props.find(qn('w:tblStyle'));original=old.get(qn('w:val')) if old is not None else None
        if old is None:old=OxmlElement('w:tblStyle');props.insert(0,old)
        name=target.tables[0].style.style_id if target.tables else original if any(s.style_id==original for s in target.styles) else 'TableGrid'
        old.set(qn('w:val'),name)


def _relationships(node,source,target):
    relmap={}
    for element in node.iter():
        for attr,value in list(element.attrib.items()):
            if not attr.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):continue
            if value not in relmap:
                rel=source.part.rels[value]
                if rel.is_external:relmap[value]=target.part.relate_to(rel.target_ref,rel.reltype,is_external=True)
                elif rel.reltype.endswith('/image'):relmap[value]=target.part.get_or_add_image(BytesIO(rel.target_part.blob))[0]
                else:raise ValueError('Unsupported linked Word object; transfer it manually in Word.')
            element.set(attr,relmap[value])
    # Drawing identifiers must be unique in the destination package.
    largest=max([int(v) for v in xp(target.element,'.//wp:docPr/@id')]+[0])
    for drawing in xp(node,'.//wp:docPr'):largest+=1;drawing.set('id',str(largest))


def _chunk(work,placement,source,target,settings,numbering):
    element=work.element(placement['source'])
    if element.get('xml'):
        node=parse_xml(element['xml'].encode())
        if element['kind']=='table' and placement.get('rows')!=element.get('rows'):
            from docx.table import Table
            table=Table(node,target._body)
            for row_index,row in enumerate(placement['rows']):
                for column,value in enumerate(row):
                    cell=table.cell(row_index,column)
                    # Do not destroy images, nested tables or merged-cell structure on text edits.
                    original=element['rows'][row_index][column]
                    if value!=original:
                        if len(cell.tables) or xp(cell._tc,'.//w:drawing | .//w:pict'):raise ValueError('This table cell contains artwork/nested tables. Edit it in Word to preserve its contents.')
                        cell.text=value
        elif element['kind'] in {'text','heading'} and placement['text']!=element['text']:_replace_text(node,placement['text'])
        _style(node,source,target,settings);_relationships(node,source,target);_numbering(node,source,target,numbering)
        return node
    if element['kind']=='table':
        rows=placement['rows'];table=target.add_table(rows=len(rows),cols=max(len(r) for r in rows))
        for ri,row in enumerate(rows):
            for ci,value in enumerate(row):table.cell(ri,ci).text=value
        node=table._tbl;node.getparent().remove(node);return node
    if element['kind']=='page' and placement['text'].strip():
        p=target.add_paragraph(placement['text'],settings['body_style']);node=p._p;node.getparent().remove(node);return node
    if element['kind'] in {'image','page'} and element.get('images'):
        p=target.add_paragraph();section=target.sections[0];width=section.page_width-section.left_margin-section.right_margin
        p.add_run().add_picture(str(work.resolve(element['images'][0])),width=width)
        node=p._p;node.getparent().remove(node);return node
    p=target.add_paragraph(placement['text'],settings['body_style']);node=p._p;node.getparent().remove(node);return node


def transfer_report(work):
    return {'revision':work.state['revision'],'title':work.state['title'],'change_summary':work.state['change_summary'],'warnings':work.state['warnings'],
            'elements':[{'id':e['id'],'kind':e['kind'],'text':e['text'],'status':work.status(e['id']),'warnings':e.get('warnings',[]),'exclusion_reason':work.state['excluded'].get(e['id'])} for e in work.elements],
            'placements':work.state['placements'],'direct_word_edits':work.state.get('authored',False)}


def build_document(work,output):
    output=Path(output).resolve()
    if output.suffix.lower()!='.docx':raise ValueError('Save the revised document as .docx.')
    protected={Path(work.state['source_original']).resolve(),Path(work.state['template_original']).resolve(),work.resolve(work.state['source_file']),work.resolve('template.docx')}
    if output in protected:raise ValueError('Choose a separate output; original documents and snapshots cannot be overwritten.')
    if work.state.get('authored'):
        data=work.resolve('authored.docx').read_bytes()
    else:
        target=Document(work.resolve('template.docx'));source=Document(work.resolve(work.state['source_file'])) if work.state['source_file'].endswith('.docx') else None
        settings=learn_styles(target);body=target.element.body;anchors={d['id']:(body[d['index']] if d['id'] not in {'start','end'} else None) for d in work.destinations};tails={};numbering={}
        for placement in work.state['placements']:
            node=_chunk(work,placement,source,target,settings,numbering);destination=placement['destination'];anchor=tails.get(destination,anchors[destination])
            if anchor is not None:anchor.addnext(node)
            elif destination=='start':body.insert(0,node)
            else:body.insert(len(body)-1 if body[-1].tag==qn('w:sectPr') else len(body),node)
            tails[destination]=node
        for paragraph in list(target.paragraphs):
            if paragraph.text.strip()=='{{steps}}':paragraph._p.getparent().remove(paragraph._p)
        project=SimpleNamespace(title=work.state['title'],version=work.state['revision'])
        fill_company_fields(target,project)
        for element in [target.element]+[sec.header._element for sec in target.sections]+[sec.footer._element for sec in target.sections]:
            for key,value in [('version',work.state['revision']),('revision',work.state['revision']),('title',work.state['title'])]:replace_element_text(element,'{{'+key+'}}',value)
        target.core_properties.title=work.state['title'];target.core_properties.comments=work.state['change_summary']
        revision=work.state['revision'].upper();target.core_properties.revision=int(revision) if revision.isdigit() and int(revision)>0 else ord(revision)-64 if len(revision)==1 and 'A'<=revision<='Z' else 1
        # Ask Word to refresh TOC and cross-references on open.
        update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');target.settings.element.append(update)
        stream=BytesIO();target.save(stream);data=stream.getvalue()
    atomic_bytes(output,data);atomic_bytes(output.with_suffix('.transfer.json'),json.dumps(transfer_report(work),ensure_ascii=False,indent=2).encode())
    return output
