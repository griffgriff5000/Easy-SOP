import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from copy import deepcopy
from pathlib import Path
from PIL import Image
from pypdf import PdfReader
from docx import Document
from PySide6.QtWidgets import QApplication
import pytest
from sop_capture.models import Project,Step,Annotation
from sop_capture.storage import ProjectStore
from sop_capture.company_information import parse_information
from sop_capture.pdf_export import export_pdf
from sop_capture.word_export import export_docx

@pytest.fixture(scope='module',autouse=True)
def app():
    return QApplication.instance() or QApplication([])

HTML='<h2>Company requirements</h2><p><b>Permit:</b> ABC-123</p><ul><li>Read the risk assessment</li></ul><table><tr><td>Prepared by</td><td>William</td></tr></table>'

def test_information_parses_rich_text_and_tables():
    blocks=parse_information(HTML)
    assert blocks[0]['level']==2
    assert blocks[1]['runs'][0]['bold']
    assert blocks[2]['list']=='bullet'
    assert blocks[3]['type']=='table'
    with pytest.raises(ValueError):parse_information('x'*500001)

def test_company_information_roundtrip_and_schema_never_downgrades(tmp_path):
    store=ProjectStore(tmp_path)
    project=Project(company_information=HTML,steps=[Step(section='Preparation',annotations=[Annotation(kind='design',points=[0,0,20,20],design={'version':1,'size':[20,20],'layers':[]})])])
    store.save(project);assert store.load().company_information==HTML
    assert project.schema_version==5
    project.company_information='';store.save(project);assert project.schema_version==5
    old=Project().to_dict();old.pop('company_information');assert Project.from_dict(old).company_information==''

def test_information_precedes_steps_in_pdf_and_word(tmp_path):
    project=Project(title='Site SOP',company_information=HTML,steps=[Step(title='Open dashboard',description='Select Dashboard.')])
    pdf=PdfReader(export_pdf(project,tmp_path,tmp_path/'procedure.pdf'))
    text=''.join(page.extract_text() for page in pdf.pages)
    assert text.index('ABC-123')<text.index('Open dashboard')
    assert 'William' in text
    document=Document(export_docx(project,tmp_path,tmp_path/'procedure.docx'))
    texts=[p.text for p in document.paragraphs]
    assert texts.index('Permit: ABC-123')<next(i for i,t in enumerate(texts) if 'Open dashboard' in t)
    assert document.tables[0].cell(0,1).text=='William'

def test_original_company_template_remains_unchanged(tmp_path):
    from sop_capture.templates import TemplateLibrary
    source=tmp_path/'company.docx';doc=Document();doc.sections[0].header.paragraphs[0].text='Company controlled document'
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Document title';table.cell(0,1).text='{{title}}'
    doc.add_paragraph('Mandatory company wording');doc.add_paragraph('{{steps}}');doc.add_paragraph('Company end matter');doc.save(source)
    original=source.read_bytes();library=TemplateLibrary(tmp_path/'library');profile=library.save('Company',{},source_docx=source)
    root=tmp_path/'project';project=Project(title='My SOP',company_information=HTML,steps=[Step(title='Open dashboard')],template=library.snapshot(profile['id'],root))
    output=Document(export_docx(project,root,tmp_path/'filled.docx'))
    assert source.read_bytes()==original
    assert output.sections[0].header.paragraphs[0].text=='Company controlled document'
    assert output.tables[0].cell(0,0).text=='Document title' and output.tables[0].cell(0,1).text=='My SOP'
    paragraphs=[p.text for p in output.paragraphs]
    assert paragraphs.index('Mandatory company wording')<paragraphs.index('Company requirements')<next(i for i,t in enumerate(paragraphs) if 'Open dashboard' in t)<paragraphs.index('Company end matter')

def test_basic_editor_autosaves_and_undo_keeps_captured_steps(tmp_path):
    from sop_capture.ui.main_window import MainWindow
    from PySide6.QtTest import QTest
    app=QApplication.instance() or QApplication([])
    window=MainWindow();store=ProjectStore(tmp_path)
    step=Step(title='Open dashboard')
    window.set_project(Project(steps=[step]),store)
    window.document_fields['title'].setText('Company SOP')
    window.company_editor.text.setPlainText('Company approval required')
    window.flush_save()
    loaded=store.load()
    assert loaded.title=='Company SOP' and 'Company approval required' in loaded.company_information
    assert loaded.schema_version==5
    assert loaded.steps[0].id==step.id
    window.undo()
    assert not window.project.company_information
    assert window.project.steps[0].id==step.id
    window.redo();assert 'Company approval required' in window.project.company_information
    window.company_editor.text.setFocus();window.company_editor.text.moveCursor(window.company_editor.text.textCursor().MoveOperation.End)
    QTest.keyClicks(window.company_editor.text,' today')
    assert 'today' in window.project.company_information
    window.close()


def test_italic_underline_and_long_table_export(tmp_path):
    html='<p><i>Italic requirement</i> and <u>underlined requirement</u></p><table><tr><td>'+('Safety requirements. '*300)+'</td><td>Approval</td></tr></table>'
    project=Project(company_information=html,steps=[Step(title='Open dashboard')])
    output=export_pdf(project,tmp_path,tmp_path/'formatted.pdf')
    assert len(PdfReader(output).pages)>2
    document=Document(export_docx(project,tmp_path,tmp_path/'formatted.docx'))
    assert any(run.italic for p in document.paragraphs for run in p.runs)
    assert any(run.underline for p in document.paragraphs for run in p.runs)


def test_editor_table_heading_and_paste_are_text_only():
    from sop_capture.ui.company_editor import CompanyInformationEditor
    from PySide6.QtCore import QMimeData
    app=QApplication.instance() or QApplication([])
    editor=CompanyInformationEditor();editor.text.setPlainText('Company requirements');editor.set_heading(2)
    assert parse_information(editor.content())[0]['level']==2
    editor.text.clear();editor.insert_table(2,2)
    editor.text.textCursor().insertText('Department')
    assert parse_information(editor.content())[0]['type']=='table'
    source=QMimeData();source.setHtml('<p><b>Approval</b><img src="file:///private.png"/></p>')
    editor.text.clear();editor.text.insertFromMimeData(source)
    assert '<img' not in editor.content()
    assert parse_information(editor.content())[0]['runs'][0]['bold']
    editor.close()

def test_table_inline_formatting_stays_in_one_paragraph(tmp_path):
    from sop_capture.ui.company_editor import CompanyInformationEditor
    from PySide6.QtCore import QMimeData
    app=QApplication.instance() or QApplication([])
    html='<table><tr><td>Permit <b>ABC</b> required</td></tr></table>'
    blocks=parse_information(html)
    assert len(blocks[0]['rows'][0][0])==1
    assert ''.join(run['text'] for run in blocks[0]['rows'][0][0][0]['runs'])=='Permit ABC required'
    editor=CompanyInformationEditor();source=QMimeData();source.setHtml(html);editor.text.insertFromMimeData(source)
    assert 'Permit ABC required' in editor.text.toPlainText()
    output=Document(export_docx(Project(company_information=html),tmp_path,tmp_path/'inline.docx'))
    assert output.tables[0].cell(0,0).text=='Permit ABC required'
    editor.close()

def test_pdf_embeds_actual_bold_and_italic_faces(tmp_path):
    project=Project(company_information='<p><b>Bold requirement</b></p><p><i>Italic requirement</i></p>')
    pdf=PdfReader(export_pdf(project,tmp_path,tmp_path/'faces.pdf'))
    used={}
    def visit(text,cm,tm,font,size):
        if font:
            for label in ('Bold requirement','Italic requirement'):
                if label in text:used[label]=str(font['/BaseFont'])
    for page in pdf.pages:page.extract_text(visitor_text=visit)
    assert 'Bold' in used['Bold requirement']
    assert any(name in used['Italic requirement'] for name in ('Oblique','Italic'))
