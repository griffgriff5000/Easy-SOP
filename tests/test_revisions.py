from io import BytesIO
from pathlib import Path
import pytest
from PIL import Image
from docx import Document

from sop_capture.revisions import RevisionWorkspace


def documents(tmp_path):
    old=tmp_path/'old.docx';template=tmp_path/'template.docx'
    doc=Document();doc.add_heading('Responsibilities',1);doc.add_paragraph('The supervisor checks the equipment.')
    table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='Role';table.cell(0,1).text='Duty';table.cell(1,0).text='Supervisor';table.cell(1,1).text='Check'
    picture=BytesIO();Image.new('RGB',(80,40),'red').save(picture,format='PNG');picture.seek(0);doc.add_picture(picture);doc.save(old)
    doc=Document();doc.sections[0].header.paragraphs[0].text='Company header';doc.add_paragraph('Company cover');doc.add_heading('Roles and responsibilities',1);doc.add_paragraph('New required information')
    doc.add_heading('Procedure',1);doc.save(template)
    return old,template


def test_workspace_approval_reopen_undo_and_originals(tmp_path):
    old,template=documents(tmp_path);original=(old.read_bytes(),template.read_bytes())
    work=RevisionWorkspace.create(tmp_path/'revision',old,template,'B')
    assert {'heading','text','table','image'}<={e['kind'] for e in work.elements}
    text=next(e for e in work.elements if e['kind']=='text');destination=next(d for d in work.destinations if 'Roles' in d['label'])
    work.suggest(text['id'],destination['id'],'Matching responsibilities')
    assert work.state['placements']==[] and work.status(text['id'])=='Suggested'
    work.approve(text['id'],destination['id'],text='Updated instruction')
    assert work.status(text['id'])=='Approved'
    reopened=RevisionWorkspace.open(work.root);assert reopened.state['placements'][0]['text']=='Updated instruction'
    work.undo();assert not work.state['placements'];work.redo();assert len(work.state['placements'])==1
    with pytest.raises(ValueError):work.approve(text['id'],'invented')
    assert original==(old.read_bytes(),template.read_bytes())


def test_export_preserves_template_and_approved_table_image(tmp_path):
    from sop_capture.revision_export import build_document
    old,template=documents(tmp_path);work=RevisionWorkspace.create(tmp_path/'revision',old,template,'B')
    dest=work.destinations[-1]['id']
    table=next(e for e in work.elements if e['kind']=='table');image=next(e for e in work.elements if e['kind']=='image')
    rows=[row[:] for row in table['rows']];rows[1][1]='Revised check';work.approve(table['id'],dest,rows=rows);work.approve(image['id'],dest)
    output=build_document(work,tmp_path/'Rev B.docx');doc=Document(output)
    assert doc.sections[0].header.paragraphs[0].text=='Company header'
    assert 'Company cover' in [p.text for p in doc.paragraphs]
    assert 'The supervisor checks the equipment.' not in [p.text for p in doc.paragraphs]
    assert doc.tables[0].cell(1,1).text=='Revised check' and len(doc.inline_shapes)==1
    assert doc.core_properties.revision==2
    assert (tmp_path/'Rev B.transfer.json').exists()
    with pytest.raises(ValueError):build_document(work,old)


def test_pdf_table_and_scan_are_reviewable(tmp_path):
    from reportlab.pdfgen.canvas import Canvas
    from sop_capture.revision_import import inspect_source
    pdf=tmp_path/'old.pdf';c=Canvas(str(pdf));c.drawString(50,780,'Procedure Rev A')
    for x in (50,200,350):c.line(x,600,x,700)
    for y in (600,650,700):c.line(50,y,350,y)
    c.drawString(60,675,'Role');c.drawString(210,675,'Duty');c.drawString(60,625,'Supervisor');c.drawString(210,625,'Check')
    c.showPage();c.rect(40,400,300,100,fill=1);c.save()
    items,warnings=inspect_source(pdf,tmp_path/'assets')
    assert any(e['kind']=='table' and e['rows'][1][0]=='Supervisor' for e in items)
    assert any(e['kind']=='page' and e['page']==2 for e in items)
    assert any('scan' in w.lower() for w in warnings)


def test_assistance_is_proposal_only_and_rejects_invalid_destinations(tmp_path):
    from sop_capture.revision_assistant import suggest_placement,validate_proposal
    old,template=documents(tmp_path);work=RevisionWorkspace.create(tmp_path/'revision',old,template,'B')
    element=next(e for e in work.elements if e['kind']=='table')
    proposal=suggest_placement(element,work.destinations,'Move below roles and responsibilities')
    assert 'responsibilities' in next(d['label'].lower() for d in work.destinations if d['id']==proposal['destination'])
    assert not work.state['placements']
    with pytest.raises(ValueError):validate_proposal({'destination':'invented','reason':'x'},work.destinations)


def test_source_list_numbering_survives_company_style_mapping(tmp_path):
    from sop_capture.revision_export import build_document
    from sop_capture.company_fields import xp
    old,template=documents(tmp_path);doc=Document(old);doc.add_paragraph('First numbered action',style='List Number');doc.add_paragraph('Second numbered action',style='List Number');doc.save(old)
    work=RevisionWorkspace.create(tmp_path/'revision',old,template,'B')
    for e in work.elements:
        if 'numbered action' in e['text']:work.approve(e['id'],'end')
    result=Document(build_document(work,tmp_path/'out.docx'))
    paragraphs=[p for p in result.paragraphs if 'numbered action' in p.text]
    assert all(xp(p._p,'./w:pPr/w:numPr/w:numId') for p in paragraphs)


def test_scan_manual_transcription_exports_as_text(tmp_path):
    from sop_capture.revision_export import build_document
    from reportlab.pdfgen.canvas import Canvas
    old,template=documents(tmp_path);pdf=tmp_path/'scan.pdf';c=Canvas(str(pdf));c.rect(50,500,200,150,fill=1);c.save()
    work=RevisionWorkspace.create(tmp_path/'revision',pdf,template,'B');page=next(e for e in work.elements if e['kind']=='page')
    work.approve(page['id'],'end',text='Manually verified transcription')
    result=Document(build_document(work,tmp_path/'transcribed.docx'))
    assert 'Manually verified transcription' in [p.text for p in result.paragraphs] and not result.inline_shapes


def test_unsupported_equations_and_notes_are_never_silently_lost(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    old,template=documents(tmp_path);doc=Document(old)
    p=doc.add_paragraph();math=OxmlElement('m:oMath');run=OxmlElement('m:r');text=OxmlElement('m:t');text.text='x=2';run.append(text);math.append(run);p._p.append(math)
    p=doc.add_paragraph('Contains a note');ref=OxmlElement('w:footnoteReference');ref.set(qn('w:id'),'1');p.add_run()._r.append(ref);doc.save(old)
    work=RevisionWorkspace.create(tmp_path/'revision',old,template,'B')
    equations=[e for e in work.elements if 'x=2' in e['text']];notes=[e for e in work.elements if 'Contains a note' in e['text']]
    assert len(equations)==1 and equations[0]['blocked'] and notes[0]['blocked']
    with pytest.raises(ValueError):work.approve(notes[0]['id'],'end')


def test_revision_import_rejects_active_word_fields(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    old,template=documents(tmp_path);doc=Document(template);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'INCLUDETEXT "https://example.invalid/private"');doc.add_paragraph()._p.append(field);doc.save(template)
    with pytest.raises(ValueError,match='active Word fields'):RevisionWorkspace.create(tmp_path/'revision',old,template,'B')


def test_local_assistant_refuses_cloud_and_invalid_responses(monkeypatch):
    from sop_capture.revision_assistant import LocalAssistant
    assistant=LocalAssistant();calls=[]
    def request(route,data=None,timeout=60):
        calls.append(route)
        if route=='tags':return {'models':[{'name':'local:small'},{'name':'remote:cloud'},{'name':'hidden','remote_host':'https://remote'}]}
        if route=='show':return {'remote_model':'remote'}
        raise AssertionError('Inference must not run')
    monkeypatch.setattr(assistant,'request',request)
    assert assistant.models()==['local:small']
    with pytest.raises(ValueError,match='Remote models'):assistant.propose('local:small',{'kind':'text','text':'private'},[{'id':'end','label':'End'}],'')
    assert 'chat' not in calls


def test_revision_details_do_not_invent_company_metadata(tmp_path):
    from sop_capture.revision_export import build_document
    old,template=documents(tmp_path);doc=Document(template);table=doc.add_table(rows=2,cols=2)
    table.cell(0,0).text='Reference';table.cell(0,1).text='COMPANY-REF-007';table.cell(1,0).text='Rev';table.cell(1,1).text='A';doc.save(template)
    work=RevisionWorkspace.create(tmp_path/'revision',old,template,'B');result=Document(build_document(work,tmp_path/'out.docx'))
    assert result.tables[0].cell(0,1).text=='COMPANY-REF-007'
    assert result.tables[0].cell(1,1).text=='B'


def test_manual_transfer_can_be_accounted_for_without_rebuilding_word(tmp_path):
    from sop_capture.revision_export import build_document
    old,template=documents(tmp_path);work=RevisionWorkspace.create(tmp_path/'revision',old,template,'B');element=work.elements[0]
    edited=template.read_bytes();(work.root/'authored.docx').write_bytes(edited);work.state['authored']=True;work.save()
    work.exclude(element['id'],'Transferred manually in Word')
    assert work.status(element['id'])=='Excluded'
    assert build_document(work,tmp_path/'manual.docx').read_bytes()==edited
