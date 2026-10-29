"""Bounded basic company-information text, shared by Word and PDF exports."""
import re
from html import escape
from lxml import etree

MAX_HTML=500_000


def parse_information(value):
    if not isinstance(value,str) or len(value)>MAX_HTML:raise ValueError('Company information is too large (maximum 500 KB of formatted text).')
    if not value.strip():return []
    try:root=etree.HTML(value,etree.HTMLParser(no_network=True,recover=True))
    except (ValueError,etree.XMLSyntaxError) as exc:raise ValueError('Invalid company information.') from exc
    if root is None:return []
    body=root.find('body')
    if body is None:return []
    count=[0]
    def runs(element,flags=None):
        flags=dict(flags or {'bold':False,'italic':False,'underline':False})
        tag=element.tag.lower() if isinstance(element.tag,str) else ''
        if tag in {'script','style','head','img','object','iframe','table'}:return []
        if tag=='br':return [dict(text='\n',**flags)]
        css={k.strip().lower():v.strip().lower() for k,v in re.findall(r'([\w-]+)\s*:\s*([^;]+)',element.get('style',''))}
        flags['bold']=flags['bold'] or tag in {'b','strong','th'} or css.get('font-weight') in {'bold','600','700','800','900'}
        flags['italic']=flags['italic'] or tag in {'i','em'} or css.get('font-style')=='italic'
        flags['underline']=flags['underline'] or tag=='u' or 'underline' in css.get('text-decoration','')
        result=[]
        if element.text:result.append(dict(text=element.text,**flags))
        for child in element:
            if child.tag=='br':result.append(dict(text='\n',**flags))
            else:result.extend(runs(child,flags))
            if child.tail:result.append(dict(text=child.tail,**flags))
        return result
    def paragraph(element,list_kind=None,number=1):
        tag=element.tag.lower();alignment=element.get('align','left').lower()
        if alignment not in {'left','center','right','justify'}:alignment='left'
        return {'type':'paragraph','level':int(tag[1]) if tag in {'h1','h2','h3'} else 0,'align':alignment,'list':list_kind,'number':number,'runs':runs(element)}
    def walk(container):
        result=[];pending=[]
        base={'bold':container.tag=='th','italic':False,'underline':False}
        if container.text:pending.append(dict(text=container.text,**base))
        def flush():
            if any(run['text'].strip() for run in pending):
                result.append({'type':'paragraph','level':0,'align':'left','list':None,'number':1,'runs':list(pending)})
            pending.clear()
        for element in container:
            if not isinstance(element.tag,str):continue
            tag=element.tag.lower()
            if tag in {'p','h1','h2','h3','pre','ul','ol','table','div','section','article','td','th'}:flush()
            if tag in {'p','h1','h2','h3','pre'}:
                block=paragraph(element)
                if any(run['text'].strip() for run in block['runs']):result.append(block)
            elif tag in {'ul','ol'}:
                number=1
                for item in element:
                    if item.tag!='li':continue
                    children=[child for child in item if child.tag in {'p','h1','h2','h3'}]
                    if children:
                        for child in children:result.append(paragraph(child,'bullet' if tag=='ul' else 'number',number));number+=1
                    else:result.append(paragraph(item,'bullet' if tag=='ul' else 'number',number));number+=1
            elif tag=='table':
                if element.xpath('.//table'):raise ValueError('Nested tables are not supported in company information.')
                rows=[];occupied=set()
                for row_number,row in enumerate(element.xpath('.//tr')):
                    if row_number>=50:raise ValueError('Company tables support up to 50 rows.')
                    cells=[];column=0
                    for cell in row:
                        if cell.tag not in {'td','th'}:continue
                        while (row_number,column) in occupied:cells.append([]);column+=1
                        try:colspan=int(cell.get('colspan','1'));rowspan=int(cell.get('rowspan','1'))
                        except ValueError:raise ValueError('Invalid table cell span.')
                        if not 1<=colspan<=8 or not 1<=rowspan<=50 or column+colspan>8:raise ValueError('Company tables support up to 8 columns.')
                        cell_blocks=walk(cell)
                        cells.append(cell_blocks)
                        for extra in range(1,colspan):cells.append([])
                        for rr in range(row_number+1,row_number+rowspan):
                            for cc in range(column,column+colspan):occupied.add((rr,cc))
                        column+=colspan
                    if cells:rows.append(cells)
                if rows:
                    width=max(len(row) for row in rows)
                    rows=[row+[[] for _ in range(width-len(row))] for row in rows]
                    result.append({'type':'table','rows':rows})
            elif tag in {'div','section','article','td','th'}:result.extend(walk(element))
            elif tag not in {'script','style','img','object','iframe','head'}:pending.extend(runs(element,base))
            if element.tail:pending.append(dict(text=element.tail,**base))
        flush()
        count[0]+=len(result)
        if count[0]>2000:raise ValueError('Company information has too many paragraphs or table cells.')
        return result
    return walk(body)


def inline_html(block):
    parts=[]
    for run in block['runs']:
        text=escape(run['text']).replace('\n','<br/>')
        for key,tag in [('bold','b'),('italic','i'),('underline','u')]:
            if run[key]:text=f'<{tag}>{text}</{tag}>'
        parts.append(text)
    prefix='• ' if block['list']=='bullet' else str(block['number'])+'. ' if block['list']=='number' else ''
    return escape(prefix)+''.join(parts)


def information_html(blocks):
    """Trusted paste subset; never carry external images, links or embedded content."""
    output=[]
    for block in blocks:
        if block['type']=='table':
            output.append('<table border="1">'+''.join('<tr>'+''.join('<td>'+information_html(cell)+'</td>' for cell in row)+'</tr>' for row in block['rows'])+'</table>')
        else:
            tag='h'+str(block['level']) if block['level'] else 'p'
            output.append(f'<{tag} align="{block["align"]}">'+inline_html(block)+f'</{tag}>')
    return ''.join(output)


def word_information(document,value,body_style,heading_style,generated):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    alignments={'left':WD_ALIGN_PARAGRAPH.LEFT,'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT,'justify':WD_ALIGN_PARAGRAPH.JUSTIFY}
    def fill(paragraph,block):
        paragraph.style=heading_style if block['level'] else body_style
        paragraph.alignment=alignments[block['align']]
        paragraph.paragraph_format.keep_with_next=bool(block['level'])
        prefix='• ' if block['list']=='bullet' else str(block['number'])+'. ' if block['list']=='number' else ''
        if prefix:paragraph.add_run(prefix)
        for segment in block['runs']:
            run=paragraph.add_run(segment['text'])
            if segment['bold']:run.bold=True
            if segment['italic']:run.italic=True
            if segment['underline']:run.underline=True
    for block in parse_information(value):
        if block['type']=='paragraph':
            paragraph=document.add_paragraph();fill(paragraph,block);generated.append(paragraph._p)
        else:
            table=document.add_table(rows=len(block['rows']),cols=len(block['rows'][0]))
            borders=OxmlElement('w:tblBorders')
            for side in ('top','left','bottom','right','insideH','insideV'):
                border=OxmlElement('w:'+side);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'B8C5CF');borders.append(border)
            table._tbl.tblPr.append(borders)
            for ri,row in enumerate(block['rows']):
                for ci,cell_blocks in enumerate(row):
                    cell=table.cell(ri,ci)
                    for pi,paragraph_block in enumerate(cell_blocks):
                        fill(cell.paragraphs[0] if pi==0 else cell.add_paragraph(),paragraph_block)
            generated.append(table._tbl)


def pdf_information(value,body_style,heading_style,width):
    from reportlab.platypus import Paragraph,Table,TableStyle,Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT,TA_CENTER,TA_RIGHT,TA_JUSTIFY
    from reportlab.lib import colors
    alignments={'left':TA_LEFT,'center':TA_CENTER,'right':TA_RIGHT,'justify':TA_JUSTIFY}
    def paragraph(block,cell=False):
        base=heading_style if block['level'] and not cell else body_style
        style=ParagraphStyle('CompanyInformation',parent=base,alignment=alignments[block['align']],keepWithNext=bool(block['level'] and not cell),spaceAfter=4 if cell else 8)
        return Paragraph(inline_html(block),style)
    output=[]
    for block in parse_information(value):
        if block['type']=='paragraph':output.append(paragraph(block))
        else:
            rows=[[[paragraph(p,True) for p in cell] or [Paragraph('',body_style)] for cell in row] for row in block['rows']]
            table=Table(rows,colWidths=[width/len(rows[0])]*len(rows[0]),hAlign='LEFT',splitByRow=1,splitInRow=1)
            table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.5,colors.HexColor('#b8c5cf')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
            output.extend([table,Spacer(1,8)])
    return output
