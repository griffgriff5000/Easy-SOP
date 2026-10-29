from io import BytesIO
from docx import Document
from docx.enum.section import WD_SECTION_START
from sop_capture.word_template import prepare_template
from sop_capture.templates import TemplateLibrary
from sop_capture.models import Project,Step
from sop_capture.word_export import export_docx


def structured_source(path):
    d=Document();d.add_paragraph('COMPANY COVER','Title');d.add_paragraph('<Insert Document Title >')
    d.add_section(WD_SECTION_START.NEW_PAGE)
    d.add_heading('Document Details',1);t=d.add_table(rows=2,cols=2)
    t.cell(0,0).text='Document Title';t.cell(0,1).text='<Insert Document Title>'
    t.cell(1,0).text='Reference';t.cell(1,1).text='<Document number>'
    d.add_heading('Approval Sign Off',1);t=d.add_table(rows=2,cols=2)
    t.cell(0,0).text='Prepared By';t.cell(0,1).text='Position/Title';t.cell(1,0).text='<full name>';t.cell(1,1).text='<position>'
    d.add_heading('Important company requirement',1);d.add_paragraph('Keep this company-specific requirement.')
    d.add_heading('Using this Template',1);d.add_paragraph('Remove this section before submission');d.add_paragraph('Template instructions and example screenshot')
    d.add_heading('MEH Heading 1',1);d.add_paragraph('Body Text: Lorem ipsum example text')
    d.add_heading('Reference Documents',1);t=d.add_table(rows=2,cols=2);t.cell(0,0).text='Document number';t.cell(0,1).text='Document title'
    d.save(path);return path


def test_preserve_cover_company_tables_sections_and_unknown_requirements(tmp_path):
    source=structured_source(tmp_path/'source.docx');original=source.read_bytes()
    data,settings=prepare_template(source);d=Document(BytesIO(data));text='\n'.join(p.text for p in d.paragraphs)
    assert 'COMPANY COVER' in text and 'Keep this company-specific requirement.' in text
    assert len(d.tables)==3 and len(d.sections)==2
    assert 'Using this Template' not in text and 'Lorem ipsum' not in text
    assert '{{steps}}' in text
    assert source.read_bytes()==original


def test_generated_sop_keeps_front_matter_and_fills_recognised_fields(tmp_path):
    source=structured_source(tmp_path/'source.docx');library=TemplateLibrary(tmp_path/'library')
    profile=library.save('Structured company',{},source_docx=source,prepare_word=True)
    project=Project(title='My procedure',reference='DOC-123',author='William',steps=[Step(title='Start work')],template=library.snapshot(profile['id'],tmp_path/'project'))
    d=Document(export_docx(project,tmp_path/'project',tmp_path/'result.docx'))
    assert len(d.tables)==3
    assert d.tables[0].cell(0,1).text=='My procedure'
    assert d.tables[0].cell(1,1).text=='DOC-123'
    assert d.tables[1].cell(1,0).text=='William'
    texts='\n'.join(p.text for p in d.paragraphs)
    assert texts.index('COMPANY COVER')<texts.index('Start work')
    assert 'Keep this company-specific requirement.' in texts


def test_bound_property_and_repeated_controls_survive_field_edit():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from sop_capture.company_fields import update_control_binding,xp
    d=Document();d.core_properties.content_status='OLD-REF';controls=[]
    for i in range(2):
        control=OxmlElement('w:sdt');properties=OxmlElement('w:sdtPr');binding=OxmlElement('w:dataBinding')
        binding.set(qn('w:storeItemID'),'{core-properties}');binding.set(qn('w:xpath'),'/cp:coreProperties/cp:contentStatus')
        binding.set(qn('w:prefixMappings'),"xmlns:cp='http://schemas.openxmlformats.org/package/2006/metadata/core-properties'")
        properties.append(binding);control.append(properties);content=OxmlElement('w:sdtContent')
        paragraph=OxmlElement('w:p');run=OxmlElement('w:r');node=OxmlElement('w:t');node.text='NEW-REF' if i==0 else 'OLD-REF'
        run.append(node);paragraph.append(run);content.append(paragraph);control.append(content);d.element.body.insert(0,control);controls.append(control)
    update_control_binding(d,controls[0])
    assert d.core_properties.content_status=='NEW-REF'
    assert ''.join(xp(controls[1],'.//w:t/text()'))=='NEW-REF'
    assert xp(controls[1],'./w:sdtPr/w:dataBinding')


def test_company_field_assistant_labels_values_and_routes_metadata(tmp_path):
    from sop_capture.ui.template_text_editor import TemplateTextEditor
    source=structured_source(tmp_path/'source.docx');data,settings=prepare_template(source);source.write_bytes(data)
    editor=TemplateTextEditor(fields_only=True);editor.set_source(source,{})
    row=next(r for r in editor.records if r['canonical']=='title')
    assert 'Document Details' in row['field_label'] and 'Document Title' in row['field_label']
    index=next(i for i,r in enumerate(editor.records) if '<position>'==r['text'])
    editor.picker.setCurrentIndex(index);editor.text.setPlainText('Supervisor')
    assert editor.edits['paragraphs'][editor.records[index]['key']]['text']=='Supervisor'
    editor.close();editor.deleteLater()


def test_custom_outline_heading_ends_guidance_without_losing_requirements(tmp_path):
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d=Document();style=d.styles.add_style('Company Section',WD_STYLE_TYPE.PARAGRAPH)
    outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),'0');style.element.get_or_add_pPr().append(outline)
    d.add_paragraph('Using this Template','Company Section');d.add_paragraph('Delete this guidance')
    d.add_paragraph('Mandatory company requirement','Company Section');d.add_paragraph('Retain this substantive requirement')
    path=tmp_path/'custom.docx';d.save(path);data,_=prepare_template(path)
    result='\n'.join(p.text for p in Document(BytesIO(data)).paragraphs)
    assert 'Retain this substantive requirement' in result and 'Delete this guidance' not in result


def test_inline_title_control_and_vertical_prepared_by_form_are_filled(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from sop_capture.company_fields import fill_form,xp
    d=Document();paragraph=d.add_paragraph();control=OxmlElement('w:sdt');properties=OxmlElement('w:sdtPr')
    alias=OxmlElement('w:alias');alias.set(qn('w:val'),'Title');properties.append(alias);control.append(properties)
    content=OxmlElement('w:sdtContent');run=OxmlElement('w:r');node=OxmlElement('w:t');node.text='OLD TITLE';run.append(node);content.append(run);control.append(content);paragraph._p.append(control)
    table=d.add_table(rows=2,cols=2);table.cell(0,0).text='Prepared By';table.cell(0,1).text='<name>';table.cell(1,0).text='Reference';table.cell(1,1).text='<ref>'
    fill_form(d,Project(title='NEW TITLE',author='William',reference='REF-456'))
    assert ''.join(xp(control,'.//w:t/text()'))=='NEW TITLE'
    assert table.cell(0,1).text=='William' and table.cell(1,1).text=='REF-456'


def test_inline_reference_cell_updates_its_backing_property():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from sop_capture.company_fields import fill_form
    d=Document();d.core_properties.content_status='OLD-REF'
    table=d.add_table(rows=1,cols=2);table.cell(0,0).text='Reference'
    paragraph=table.cell(0,1).paragraphs[0]
    control=OxmlElement('w:sdt');properties=OxmlElement('w:sdtPr');binding=OxmlElement('w:dataBinding')
    binding.set(qn('w:storeItemID'),'{core-properties}');binding.set(qn('w:xpath'),'/cp:coreProperties/cp:contentStatus')
    binding.set(qn('w:prefixMappings'),"xmlns:cp='http://schemas.openxmlformats.org/package/2006/metadata/core-properties'")
    properties.append(binding);control.append(properties);content=OxmlElement('w:sdtContent')
    run=OxmlElement('w:r');node=OxmlElement('w:t');node.text='OLD-REF';run.append(node);content.append(run);control.append(content);paragraph._p.append(control)
    fill_form(d,Project(reference='NEW-REF'))
    assert d.core_properties.content_status=='NEW-REF'


def test_explicit_steps_marker_survives_guidance_removal_without_duplicate_heading(tmp_path):
    source=tmp_path/'source.docx';d=Document();d.add_heading('Company cover',0);d.add_heading('Procedure',1);d.add_paragraph('{{steps}}')
    d.add_heading('Using this Template',1);d.add_paragraph('Delete this guidance');d.save(source)
    library=TemplateLibrary(tmp_path/'library');profile=library.save('Existing marker',{},source_docx=source,prepare_word=True)
    project=Project(steps=[Step(title='Do the task')],template=library.snapshot(profile['id'],tmp_path/'project'))
    result=Document(export_docx(project,tmp_path/'project',tmp_path/'result.docx'))
    texts=[p.text for p in result.paragraphs]
    assert texts.count('Procedure')==1
    assert any('Do the task' in text for text in texts) and 'Delete this guidance' not in texts


def test_sample_text_in_one_cell_does_not_clear_the_rest_of_a_company_table(tmp_path):
    d=Document();table=d.add_table(rows=2,cols=2)
    table.cell(0,0).text='Important label';table.cell(0,1).text='Retain this company value'
    table.cell(1,0).text='Description';table.cell(1,1).text='Lorem ipsum sample'
    path=tmp_path/'form.docx';d.save(path);data,_=prepare_template(path);result=Document(BytesIO(data))
    assert result.tables[0].cell(0,0).text=='Important label'
    assert result.tables[0].cell(0,1).text=='Retain this company value'
    assert result.tables[0].cell(1,1).text==''
