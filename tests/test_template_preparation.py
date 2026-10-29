from io import BytesIO
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt,RGBColor,Mm
from PIL import Image
from sop_capture.templates import TemplateLibrary
from sop_capture.models import Project,Step
from sop_capture.storage import ProjectStore
from sop_capture.word_export import export_docx


def company_template(path):
    document=Document()
    for name,size in [('ACME_Body Text',10),('ACME_Heading 1',18),('ACME_Heading 2',16)]:
        style=document.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);style.font.name='Arial';style.font.size=Pt(size)
        if 'Heading' in name:style.font.color.rgb=RGBColor.from_string('79C000')
    document.sections[0].left_margin=Mm(24)
    header=document.sections[0].header.paragraphs[0];header.text='CONTROLLED | <Insert Document Title>'
    logo=BytesIO();Image.new('RGB',(32,16),'green').save(logo,'PNG');logo.seek(0);header.add_run().add_picture(logo)
    document.sections[0].footer.paragraphs[0].text='COMPANY TEMPLATE ID | UNCONTROLLED WHEN PRINTED'
    document.add_paragraph('Using this Template','ACME_Heading 1')
    document.add_paragraph('Remove this section before submission','ACME_Body Text')
    document.add_paragraph('Example heading','ACME_Heading 2');document.add_paragraph('Lorem ipsum example instructions','ACME_Body Text')
    document.add_table(rows=1,cols=2).cell(0,0).text='Sample table to remove'
    document.save(path);return path


def test_upload_clears_guidance_and_uses_corporate_styles(tmp_path):
    from sop_capture.ui.templates_page import TemplatesPage
    source=company_template(tmp_path/'company.docx');original=source.read_bytes();library=TemplateLibrary(tmp_path/'library')
    page=TemplatesPage(library,lambda identity:None,lambda:None)
    profile=page.import_path(str(source),'Company')
    assert profile.get('prepared') is True
    assert profile['settings']['body_style']=='ACME_Body Text'
    assert profile['settings']['heading_style']=='ACME_Heading 2'
    assert profile['settings']['title_style']=='ACME_Heading 1'
    root=tmp_path/'project';project=Project(title='Set to work',steps=[Step(title='Open dashboard',description='Choose your job')],template=library.snapshot(profile['id'],root))
    doc=Document(export_docx(project,root,tmp_path/'result.docx'))
    text='\n'.join(p.text for p in doc.paragraphs)
    assert 'Using this Template' not in text and 'Lorem ipsum' not in text
    assert not doc.tables
    assert next(p for p in doc.paragraphs if p.text=='Choose your job').style.name=='ACME_Body Text'
    assert next(p for p in doc.paragraphs if 'Open dashboard' in p.text).style.name=='ACME_Heading 2'
    assert next(p for p in doc.paragraphs if p.text=='Set to work').style.name=='ACME_Heading 1'
    assert 'Set to work' in doc.sections[0].header.paragraphs[0].text
    assert 'COMPANY TEMPLATE ID' in doc.sections[0].footer.paragraphs[0].text
    assert len(doc.sections[0].header._element.xpath('.//w:drawing'))==1
    assert round(doc.sections[0].left_margin.mm)==24
    assert source.read_bytes()==original
    assert 'Remove this section' in '\n'.join(p.text for p in Document(source).paragraphs)
    updated=page.save_changes();assert updated['prepared'] is True
    page.close();page.deleteLater()


def test_style_numbering_is_not_duplicated_and_title_is_not_numbered(tmp_path):
    source=company_template(tmp_path/'company.docx');document=Document(source)
    for name in ('ACME_Heading 1','ACME_Heading 2'):
        num=OxmlElement('w:numPr');identifier=OxmlElement('w:numId');identifier.set(qn('w:val'),'1');num.append(identifier);document.styles[name].element.get_or_add_pPr().append(num)
    document.save(source)
    library=TemplateLibrary(tmp_path/'library');profile=library.save('Company',{},source_docx=source,prepare_word=True)
    root=tmp_path/'project';project=Project(title='SOP title',steps=[Step(title='Click here')],template=library.snapshot(profile['id'],root))
    doc=Document(export_docx(project,root,tmp_path/'output.docx'))
    heading=next(p for p in doc.paragraphs if 'Click here' in p.text)
    assert heading.text=='Click here'
    title=next(p for p in doc.paragraphs if p.text=='SOP title')
    assert title._p.xpath('./w:pPr/w:numPr/w:numId/@w:val')==['0']


def test_header_prompt_split_over_runs_preserves_artwork_in_same_run(tmp_path):
    from sop_capture.word_template import fill_company_fields
    document=Document();paragraph=document.sections[0].header.paragraphs[0]
    paragraph.add_run('Owner | <Insert ')
    run=paragraph.add_run('Document Title>');run.bold=True
    data=BytesIO();Image.new('RGB',(20,20),'blue').save(data,'PNG');data.seek(0);run.add_picture(data)
    fill_company_fields(document,Project(title='My SOP'))
    assert paragraph.text=='Owner | My SOP'
    assert len(paragraph._p.xpath('.//w:drawing'))==1


def test_preparation_preserves_company_section_layouts(tmp_path):
    from docx.enum.section import WD_SECTION_START
    from sop_capture.word_template import prepare_template
    source=company_template(tmp_path/'source.docx');document=Document(source)
    new=document.add_section(WD_SECTION_START.NEW_PAGE);new.page_width=Mm(297);new.page_height=Mm(210);new.left_margin=Mm(12)
    document.add_paragraph('Landscape example');document.save(source)
    data,settings=prepare_template(source);clean=Document(BytesIO(data))
    assert len(clean.sections)==2 and not clean.tables
    assert '{{steps}}' in [p.text for p in clean.paragraphs]
    assert 'Landscape example' not in [p.text for p in clean.paragraphs]
    assert round(clean.sections[0].page_width.mm)==216  # synthetic source's Letter page
    assert round(clean.sections[0].left_margin.mm)==24
    assert round(clean.sections[1].left_margin.mm)==12
    assert 'CONTROLLED' in clean.sections[0].header.paragraphs[0].text


def test_unfamiliar_style_names_use_body_usage_and_outline_levels(tmp_path):
    from sop_capture.word_template import learn_styles
    document=Document()
    for name in ('Company instruction','Chapter name','Task name'):
        document.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    for name,level in [('Chapter name',0),('Task name',1)]:
        outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),str(level));document.styles[name].element.get_or_add_pPr().append(outline)
    document.add_paragraph('Some useful instructions','Company instruction');document.add_paragraph('More instructions','Company instruction')
    document.add_paragraph('Example chapter','Chapter name');document.add_paragraph('Example task','Task name')
    settings=learn_styles(document)
    assert settings['body_style']=='Company instruction'
    assert settings['title_style']=='Chapter name'
    assert settings['heading_style']=='Task name'
