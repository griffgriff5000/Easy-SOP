import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from io import BytesIO
from docx import Document
from sop_capture.models import Project,Step
from sop_capture.storage import ProjectStore
from sop_capture.authored_document import store_document,source_fingerprint,make_working_copy
from sop_capture.word_export import export_docx
import pytest


def document_bytes(text):
    document=Document();document.add_paragraph(text);data=BytesIO();document.save(data);return data.getvalue()


def test_saved_document_is_authoritative_and_immutable(tmp_path):
    store=ProjectStore(tmp_path);project=Project(title='Source title',steps=[Step(title='Source step')])
    data=document_bytes('Manually edited company document')
    project.document_file=store_document(store,data);project.document_source_key=source_fingerprint(project)
    store.save(project);assert store.load().schema_version==7
    project.title='Changed source title';project.steps.append(Step(title='New capture'))
    output=export_docx(project,tmp_path,tmp_path/'export.docx')
    assert output.read_bytes()==data
    assert Document(output).paragraphs[0].text=='Manually edited company document'
    assert store.resolve(project.document_file).read_bytes()==data


def test_working_copy_preserves_edits_until_explicit_rebuild(tmp_path):
    store=ProjectStore(tmp_path);project=Project(title='Source',steps=[Step(title='Captured step')])
    project.document_file=store_document(store,document_bytes('Edited body'));key=source_fingerprint(project);project.document_source_key=key
    first=make_working_copy(project,tmp_path,tmp_path/'working.docx')
    assert Document(first).paragraphs[0].text=='Edited body'
    make_working_copy(project,tmp_path,tmp_path/'rebuilt.docx',rebuild=True)
    assert any('Captured step' in p.text for p in Document(tmp_path/'rebuilt.docx').paragraphs)
    assert source_fingerprint(project)==key
    project.steps.append(Step());assert source_fingerprint(project)!=key


def test_authored_document_path_and_data_validation(tmp_path):
    store=ProjectStore(tmp_path)
    with pytest.raises(ValueError):store_document(store,b'not a document')
    project=Project(document_file='../outside.docx')
    with pytest.raises(ValueError):store.save(project)


def test_editor_save_failure_keeps_working_document_and_retries(tmp_path):
    from sop_capture.ui.word_editor import WordEditor
    class Session:
        closed=False
        def save(self):return True
        def close(self):self.closed=True
    saved=[];fail=[True]
    def persist(data,key):
        if fail[0]:raise OSError('Disk full')
        saved.append(data);return 'Documents/saved.docx'
    editor=WordEditor(persist);editor.path=tmp_path/'working.docx';editor.path.write_bytes(document_bytes('Unsaved direct edits'))
    session=Session();editor.session=session
    with pytest.raises(OSError,match='Disk full'):editor.close_document()
    assert editor.session is session and not session.closed
    assert editor.path.exists() and editor.saved_hash is None
    fail[0]=False;editor.save_document();assert len(saved)==1
    editor.save_document();assert len(saved)==1  # no redundant immutable copies
    editor.close_document();assert session.closed


def test_project_snapshot_failure_rolls_back_without_losing_native_work(tmp_path,monkeypatch):
    from sop_capture.ui.main_window import MainWindow
    window=MainWindow();store=ProjectStore(tmp_path);window.set_project(Project(),store)
    old=window.project.to_dict();history=list(window.history)
    save=store.save
    def fail(project):raise OSError('Cannot save project')
    monkeypatch.setattr(store,'save',fail)
    with pytest.raises(OSError):window._persist_document(document_bytes('Direct edits'),'0'*64)
    assert window.project.to_dict()==old and window.history==history
    monkeypatch.setattr(store,'save',save);window.close()


def test_pdf_uses_authored_document_even_for_standard_template(tmp_path,monkeypatch):
    from sop_capture.document_export import export_pdf
    import sop_capture.word_export as exporter
    store=ProjectStore(tmp_path);project=Project()
    data=document_bytes('User edited document');project.document_file=store_document(store,data)
    observed=[]
    def pdf(source,output):
        observed.append(source.read_bytes());output.write_bytes(b'%PDF-test');return output
    monkeypatch.setattr(exporter,'word_pdf',pdf)
    assert export_pdf(project,tmp_path,tmp_path/'output.pdf').read_bytes()==b'%PDF-test'
    assert observed==[data]


def test_switch_project_does_not_close_editor_when_save_fails(tmp_path,monkeypatch):
    from sop_capture.ui.main_window import MainWindow
    window=MainWindow();store=ProjectStore(tmp_path/'first');project=Project();window.set_project(project,store)
    def fail():raise OSError('Unsaved document')
    monkeypatch.setattr(window.word_editor,'save_document',fail)
    with pytest.raises(OSError):window.set_project(Project(),ProjectStore(tmp_path/'second'))
    assert window.project is project and window.store is store and store.lock_token
    monkeypatch.undo();window.close()


def test_native_save_as_uses_new_document_path(tmp_path):
    from sop_capture.ui.word_editor import WordEditor
    original=tmp_path/'working.docx';original.write_bytes(document_bytes('Original working document'))
    new=tmp_path/'renamed.docx';new.write_bytes(document_bytes('Edits after native Save As'))
    class Session:
        saved_path=new
        def save(self):return True
        def close(self):pass
    snapshots=[]
    editor=WordEditor(lambda data,key:snapshots.append(data) or 'Documents/snapshot.docx')
    editor.path=original;editor.session=Session();editor.save_document()
    assert snapshots==[new.read_bytes()]
    editor.close_document()


@pytest.mark.parametrize('existing_count,reused_pid',[(1,False),(0,True)])
def test_word_session_never_quits_an_existing_word_instance(tmp_path,monkeypatch,existing_count,reused_pid):
    from types import SimpleNamespace
    import pythoncom,win32com.client,win32process
    import sop_capture.word_session as module
    calls=[]
    class Probe:
        ActiveWindow=SimpleNamespace(Hwnd=123)
        def Close(self,save):calls.append('close probe')
    class Documents:
        Count=existing_count
        def Add(self):calls.append('add probe');return Probe()
    word=SimpleNamespace(Documents=Documents(),Quit=lambda:calls.append('quit'))
    monkeypatch.setattr(module.subprocess,'run',lambda *a,**k:SimpleNamespace(stdout='"WINWORD.EXE","42","Console","1","10 K"'))
    monkeypatch.setattr(win32com.client,'DispatchEx',lambda name:word)
    monkeypatch.setattr(win32process,'GetWindowThreadProcessId',lambda hwnd:(1,42 if reused_pid else 43))
    monkeypatch.setattr(pythoncom,'CoInitialize',lambda:None);monkeypatch.setattr(pythoncom,'CoUninitialize',lambda:calls.append('uninitialize'))
    with pytest.raises(RuntimeError,match='reused'):module.WordSession(tmp_path/'unused.docx')
    assert 'quit' not in calls and calls[-1]=='uninitialize'
    assert ('add probe' in calls)==(existing_count==0)
