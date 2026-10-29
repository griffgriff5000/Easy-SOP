"""Prepare the supplied company template and verify actual Word/PDF output."""
from pathlib import Path
import sys,tempfile,logging
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from PySide6.QtWidgets import QApplication
from PySide6.QtPdf import QPdfDocument
from PySide6.QtCore import QSize
from docx import Document
from PIL import Image
from pypdf import PdfReader
from sop_capture.templates import TemplateLibrary
from sop_capture.storage import ProjectStore
from sop_capture.models import Project,Step
from sop_capture.ui.templates_page import TemplatesPage
from sop_capture.word_export import export_docx
from sop_capture.document_export import export_pdf


def main():
    logging.basicConfig(level=logging.INFO)
    app=QApplication([])
    source=Path(sys.argv[1]);original=source.read_bytes()
    output=Path('.superpowers/prepared-template');output.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='sop-template-qa-') as temporary:
        root=Path(temporary);library=TemplateLibrary(root/'library');page=TemplatesPage(library,lambda identity:None,lambda:None)
        profile=page.import_path(str(source),'MEH company template')
        print('Imported company template; structured:',profile.get('structured'),flush=True)
        output=output/('structured' if profile.get('structured') else 'simple');output.mkdir(parents=True,exist_ok=True)
        store=ProjectStore(root/'project')
        project=Project(title='Set to work',reference='SOP-QA-001',author='Example author',reviewer='Example reviewer',approver='Example approver',purpose='Record the process for opening a job.',steps=[Step(title='Open the dashboard',description='Select the required project.'),Step(title='Choose the job',description='Confirm the job details before starting work.')],template=library.snapshot(profile['id'],store.root))
        project.steps[0].image=store.add_image(Image.new('RGB',(900,500),'#e7f0f5'),project.steps[0].id)
        if profile.get('structured'):
            from sop_capture.template_text import template_paragraphs,edit_snapshot
            records=template_paragraphs(store.resolve(project.template['docx']))
            team=next((row for row in records if 'Teamcenter ID' in row.get('field_label','') and row['editable']),None)
            if team:project.template_text_edits=edit_snapshot(store.resolve(project.template['docx']),{team['key']:'TC-TEST-123'})
        export_docx(project,store.root,output/'MEH SOP.docx');export_pdf(project,store.root,output/'MEH SOP.pdf')
        document=Document(output/'MEH SOP.docx');body='\n'.join(p.text for p in document.paragraphs)
        assert 'Using this Template' not in body and 'Lorem ipsum' not in body
        assert profile['settings']['body_style']=='MEH_Body Text'
        assert profile['settings']['heading_style']=='MEH_Heading 2'
        assert source.read_bytes()==original
        pdf=PdfReader(output/'MEH SOP.pdf');text='\n'.join(p.extract_text() for p in pdf.pages)
        assert 'MEH ALLIANCE' in text or 'UK PROTECT' in text
        assert 'Open the dashboard' in text and 'Using this Template' not in text
        if profile.get('structured'):
            assert len(document.tables)>=6 and len(document.sections)==2
            assert 'Document Details' in text and 'Approval Sign Off' in text
            assert 'TO USE TEMPLATE' not in text and 'Lorem ipsum' not in text
            assert 'TC-TEST-123' in text and '123456789 / N/A' not in text
            assert 'Example reviewer' in text and 'SOP-QA-001' in text
        render=QPdfDocument();render.load(str((output/'MEH SOP.pdf').resolve()))
        for i in range(render.pageCount()):
            path=output/f'page-{i+1}.png';image=render.render(i,QSize(1000,1414));image.save(str(path))
            with Image.open(path) as rendered:
                rgba=rendered.convert('RGBA');paper=Image.new('RGBA',rgba.size,'white');paper.alpha_composite(rgba);paper.convert('RGB').save(path)
        render.close();page.close();app.processEvents()
        print('PASS: supplied MEH template prepared automatically, company styles and furniture retained, sample body removed, source unchanged.')
        print('Detected styles:',profile['settings']['title_style'],profile['settings']['heading_style'],profile['settings']['body_style'])
        print('Rendered pages:',len(pdf.pages));print(text[:1800])


if __name__=='__main__':main()
