"""Actual Word embedding, direct edits, immutable save/reopen and export check."""
from pathlib import Path
import sys
import tempfile
import time
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from PySide6.QtWidgets import QApplication
from sop_capture.models import Project,Step
from sop_capture.storage import ProjectStore
from sop_capture.templates import TemplateLibrary
from sop_capture.ui.main_window import MainWindow
from sop_capture.word_export import export_docx
from sop_capture.document_export import export_pdf


def wait(app, predicate, timeout=45):
    deadline=time.monotonic()+timeout
    while not predicate() and time.monotonic()<deadline:
        app.processEvents();time.sleep(.02)
    if not predicate():raise RuntimeError('Timed out waiting for the document editor')


def main():
    app=QApplication([])
    output=Path('.superpowers/direct-editor');output.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='sop-direct-qa-') as temporary:
        root=Path(temporary);source=root/'company.docx'
        document=Document();document.sections[0].header.paragraphs[0].text='COMPANY CONTROLLED DOCUMENT'
        document.add_heading('{{title}}',0);document.add_paragraph('Remove these template instructions')
        table=document.add_table(rows=1,cols=2);table.style='Table Grid';table.cell(0,0).text='Owner';table.cell(0,1).text='Original owner'
        document.add_paragraph('{{steps}}');document.save(source);original=source.read_bytes()
        library=TemplateLibrary(root/'library');profile=library.save('Company template',{},source_docx=source)
        store=ProjectStore(root/'project');project=Project(title='Direct document editing',steps=[Step(title='Open the dashboard',description='Choose the required job.')],template=library.snapshot(profile['id'],store.root))
        window=MainWindow(template_library=library);window.set_project(project,store);window.show();window.show_page(2)
        try:
            wait(app,lambda:window.word_editor.session is not None and not window.busy)
            editor=window.word_editor
            if not editor.container:raise RuntimeError(editor.status.text())
            doc=editor.session.document
            # Native Word selection and typing, not conversion to Qt rich text.
            selection=doc.ActiveWindow.Selection;selection.EndKey(6);selection.TypeText('\rDirect edit retained in project and exports.')
            doc.Tables(1).Cell(1,2).Range.Text='William Griffin'
            instructions=doc.Content.Duplicate
            if instructions.Find.Execute(FindText='Remove these template instructions'):instructions.Delete()
            window.flush_save();reference=project.document_file
            assert reference and store.load().document_file==reference
            import win32gui
            assert win32gui.GetParent(editor.session.hwnd)!=0
            doc.Activate();editor.session.word.ScreenRefresh()
            deadline=time.monotonic()+2
            while time.monotonic()<deadline:app.processEvents();time.sleep(.02)
            from PIL import ImageGrab
            ImageGrab.grab(bbox=win32gui.GetWindowRect(int(window.winId()))).save(output/'editable-centre.png')
            saved=Document(store.resolve(reference))
            assert saved.tables[0].cell(0,1).text=='William Griffin'
            assert any('Direct edit retained' in p.text for p in saved.paragraphs)
            assert not any('Remove these template instructions' in p.text for p in saved.paragraphs)
            assert source.read_bytes()==original
            doc.SaveAs2(str(root/'native-save-as.docx'),FileFormat=12)
            doc.ActiveWindow.Selection.EndKey(6);doc.ActiveWindow.Selection.TypeText('\rNative Save As edits retained.')
            window.flush_save()
            assert any('Native Save As edits retained' in p.text for p in Document(store.resolve(project.document_file)).paragraphs)
            editor.close_document();window.open_document_editor()
            wait(app,lambda:editor.session is not None and not window.busy)
            assert 'Direct edit retained' in editor.session.document.Content.Text
            assert 'Native Save As edits retained' in editor.session.document.Content.Text
            window.flush_save()
            export_docx(project,store.root,output/'edited-document.docx')
            export_pdf(project,store.root,output/'edited-document.pdf')
            from pypdf import PdfReader
            text=''.join(p.extract_text() for p in PdfReader(output/'edited-document.pdf').pages)
            assert 'Direct edit retained' in text and 'William Griffin' in text
            assert 'Open the dashboard' in text and 'Remove these template instructions' not in text
            last=editor.session.document
            last.SaveAs2(str(root/'closed-native-save-as.docx'),FileFormat=12)
            last.ActiveWindow.Selection.EndKey(6);last.ActiveWindow.Selection.TypeText('\rSaved before closing native Word.')
            last.Save();last.Close(False)
            editor.save_document()
            assert editor.session is None
            assert any('Saved before closing native Word' in p.text for p in Document(store.resolve(project.document_file)).paragraphs)
            print('PASS: embedded Word, native typing, table edit, paragraph deletion, project save/reopen, original unchanged, DOCX/PDF exports.')
        finally:
            window.close();app.processEvents()


if __name__=='__main__':main()
