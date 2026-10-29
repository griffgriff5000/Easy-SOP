"""Side-by-side human-approved procedure revision, independent of SOP capture."""
from html import escape
from pathlib import Path
import tempfile
from docx import Document
from PySide6.QtCore import Qt,QBuffer,QIODevice,QUrl,QLockFile
from PySide6.QtGui import QImage,QTextDocument
from PySide6.QtWidgets import (QWidget,QVBoxLayout,QHBoxLayout,QLabel,QPushButton,QSplitter,QListWidget,QTextBrowser,QPlainTextEdit,QLineEdit,QComboBox,QTabWidget,QTableWidget,QTableWidgetItem,QFileDialog,QInputDialog,QMessageBox,QDialog,QFormLayout,QDialogButtonBox)
from ..revisions import RevisionWorkspace
from ..revision_export import build_document
from ..revision_assistant import suggest_placement,LocalAssistant
from ..storage import atomic_bytes
from ..company_fields import xp
from ..revision_import import element_text
from .components import page_header
from .word_editor import WordEditor


def document_html(browser,path):
    """Readable structure preview; native Word remains the authoritative page layout."""
    doc=Document(path);body=[]
    def render(node):
        if node.tag.endswith('}tbl'):
            return '<table border="1" cellspacing="0" cellpadding="5">'+''.join('<tr>'+''.join('<td>'+''.join(render(p) for p in xp(cell,'./w:p | ./w:tbl'))+'</td>' for cell in xp(row,'./w:tc'))+'</tr>' for row in xp(node,'./w:tr'))+'</table>'
        text=escape(element_text(node)).replace('\n','<br>');pictures=''
        for rid in xp(node,'.//@r:embed'):
            rel=doc.part.rels.get(str(rid))
            if rel and not rel.is_external and rel.reltype.endswith('/image'):
                image=QImage.fromData(rel.target_part.blob)
                if not image.isNull():
                    url=QUrl('revision-image-'+str(len(body))+'-'+str(rid));browser.document().addResource(QTextDocument.ResourceType.ImageResource,url,image)
                    width=min(400,image.width());pictures+=f'<br><img src="{url.toString()}" width="{width}">'
        return '<p>'+text+pictures+'</p>'
    body.append('<p><i>Structure preview — open the Word editor to check exact pages, styles and pagination.</i></p>')
    for sec in doc.sections:
        if any(p.text for p in sec.header.paragraphs):body.append('<p><i>Template header: '+escape(' '.join(p.text for p in sec.header.paragraphs))+'</i></p>');break
    for node in doc.element.body:body.append(render(node))
    return ''.join(body)


class NewRevisionDialog(QDialog):
    def __init__(self,library,parent=None):
        super().__init__(parent);self.setWindowTitle('Revise an existing document');self.resize(680,360)
        box=QVBoxLayout(self);form=QFormLayout();box.addLayout(form);self.old=QLineEdit();self.template=QLineEdit();self.folder=QLineEdit();self.revision=QLineEdit('B')
        for label,field,kind in [('Old Word or PDF',self.old,'old'),('New Word template',self.template,'template'),('New workspace folder',self.folder,'folder')]:
            row=QHBoxLayout();row.addWidget(field,1);button=QPushButton('Browse…');button.clicked.connect(lambda checked=False,f=field,k=kind:self.browse(f,k));row.addWidget(button);form.addRow(label,row)
        self.saved=QComboBox();self.saved.addItem('Or choose a saved Word template…',None)
        if library:
            for profile in library.list():
                if profile['kind']=='docx':self.saved.addItem(profile['name'],str(library.asset(profile,'docx')))
        self.saved.currentIndexChanged.connect(lambda:self.template.setText(self.saved.currentData() or self.template.text()));form.addRow('Template library',self.saved);form.addRow('New revision',self.revision)
        note=QLabel('Choose a new, unused workspace folder. Source and template copies, approvals and the working document stay there. Original files are kept unchanged.');note.setWordWrap(True);box.addWidget(note)
        buttons=QDialogButtonBox(QDialogButtonBox.StandardButton.Ok|QDialogButtonBox.StandardButton.Cancel);buttons.button(QDialogButtonBox.StandardButton.Ok).setText('Inspect documents');buttons.accepted.connect(self.accept);buttons.rejected.connect(self.reject);box.addWidget(buttons)

    def browse(self,field,kind):
        if kind=='folder':
            parent=QFileDialog.getExistingDirectory(self,'Choose parent folder')
            if parent:field.setText(str(Path(parent)/('Revision '+self.revision.text().strip())))
        else:
            path,_=QFileDialog.getOpenFileName(self,'Choose document','','Word or PDF (*.docx *.pdf)' if kind=='old' else 'Word document (*.docx)')
            if path:field.setText(path)

    def accept(self):
        old,template,folder=Path(self.old.text()),Path(self.template.text()),Path(self.folder.text())
        if not old.is_file() or old.suffix.lower() not in {'.docx','.pdf'} or not template.is_file() or template.suffix.lower()!='.docx' or not folder.is_absolute() or folder.exists() or not self.revision.text().strip():
            QMessageBox.warning(self,'Check document choices','Choose an existing Word/PDF source, a Word template, a new absolute workspace folder and a revision.');return
        super().accept()


class RevisionPage(QWidget):
    def __init__(self,library,run_task=None):
        super().__init__();self.library=library;self.work=None;self.run_task=run_task;self.pdf_document=None;self.pdf_buffer=None;self.loading=False;self.workspace_lock=None
        box=QVBoxLayout(self);box.setSpacing(10);page_header(box,'Revise document','Compare the old procedure with a new Word template. Review and approve every placement.')
        tools=QHBoxLayout();box.addLayout(tools)
        for text,callback in [('New revision',self.new),('Open revision',self.open),('Original pages',self.original_pages),('Undo',self.undo),('Redo',self.redo),('Export Word + report',self.export)]:
            button=QPushButton(text);button.clicked.connect(callback);tools.addWidget(button)
        tools.addStretch();self.location=QLabel('No revision workspace open');self.location.setWordWrap(True);box.addWidget(self.location)
        details=QHBoxLayout();self.title=QLineEdit();self.title.setPlaceholderText('Document title');details.addWidget(self.title,2);self.revision=QLineEdit();self.revision.setMaximumWidth(90);self.revision.setPlaceholderText('Rev');details.addWidget(self.revision)
        self.summary=QLineEdit();self.summary.setPlaceholderText('Revision/change summary');details.addWidget(self.summary,2);self.details_button=QPushButton('Save details');self.details_button.clicked.connect(self.save_details);details.addWidget(self.details_button);box.addLayout(details)
        split=QSplitter();box.addWidget(split,1)
        left=QWidget();l=QVBoxLayout(left);l.setContentsMargins(0,0,0,0);l.addWidget(QLabel('OLD REVISION — source remains unchanged'))
        self.sources=QListWidget();self.sources.setMaximumHeight(170);self.sources.currentRowChanged.connect(self.select_source);l.addWidget(self.sources)
        self.source_tabs=QTabWidget();self.source_view=QTextBrowser();self.source_view.setOpenExternalLinks(False);self.source_tabs.addTab(self.source_view,'Selected element')
        self.original=QTextBrowser();self.source_tabs.addTab(self.original,'Whole document');l.addWidget(self.source_tabs,1);split.addWidget(left)
        right=QWidget();r=QVBoxLayout(right);r.setContentsMargins(0,0,0,0);r.addWidget(QLabel('NEW REVISION — proposed and approved content'))
        self.right_tabs=QTabWidget();r.addWidget(self.right_tabs,1);split.addWidget(right);split.setSizes([500,650])
        review=QWidget();reviewbox=QVBoxLayout(review);self.destination=QComboBox();reviewbox.addWidget(QLabel('Place after template heading'));reviewbox.addWidget(self.destination)
        order=QWidget();orderbox=QVBoxLayout(order);orderbox.addWidget(QLabel('Approved placements — select one to return to its content.'))
        self.placements=QListWidget();self.placements.currentRowChanged.connect(self.select_placement);orderbox.addWidget(self.placements,1)
        moves=QHBoxLayout()
        for label,callback in [('Move up',lambda:self.move(-1)),('Move down',lambda:self.move(1)),('Return to unplaced',self.unplace)]:
            button=QPushButton(label);button.clicked.connect(callback);moves.addWidget(button)
        orderbox.addLayout(moves);self.editor_label=QLabel('Edit text or table cells before approval.');self.editor_label.setWordWrap(True);reviewbox.addWidget(self.editor_label)
        self.content=QTabWidget();self.text=QPlainTextEdit();self.table=QTableWidget();self.table.setMinimumHeight(140);self.content.addTab(self.text,'Text');self.content.addTab(self.table,'Table cells');reviewbox.addWidget(self.content,1)
        self.proposal=QLabel('Select an element from the old revision.');self.proposal.setWordWrap(True);reviewbox.addWidget(self.proposal)
        approval=QHBoxLayout();self.approve_button=QPushButton('Approve placement');self.approve_button.setObjectName('primary');self.approve_button.clicked.connect(self.approve);approval.addWidget(self.approve_button)
        self.exclude_button=QPushButton('Exclude with reason');self.exclude_button.clicked.connect(self.exclude);approval.addWidget(self.exclude_button);reviewbox.addLayout(approval);self.right_tabs.addTab(review,'Placement');self.right_tabs.addTab(order,'Order')
        self.new_view=QTextBrowser();self.new_view.setOpenExternalLinks(False);self.right_tabs.addTab(self.new_view,'Preview')
        final=QWidget();finalbox=QVBoxLayout(final);self.word_editor=WordEditor(self.persist_document,self);self.word_editor.rebuild_button.setText('Return to placements…');self.word_editor.rebuild_requested.connect(self.return_to_placements);self.word_editor.open_requested.connect(self.open_word);self.word_editor.problem.connect(self.error);finalbox.addWidget(self.word_editor);self.right_tabs.addTab(final,'Word')
        assistant=QHBoxLayout();self.instruction=QLineEdit();self.instruction.setPlaceholderText('Ask about selected element, e.g. Move below Responsibilities');assistant.addWidget(self.instruction,1)
        self.model=QComboBox();self.model.addItem('Local structural matching',None);self.model.setMaximumWidth(210);assistant.addWidget(self.model)
        models=QPushButton('Find local AI');models.clicked.connect(self.find_models);assistant.addWidget(models);self.ask_button=QPushButton('Suggest');self.ask_button.clicked.connect(self.ask);assistant.addWidget(self.ask_button);box.addLayout(assistant)
        self.status=QLabel('Offline structural matching is built in. Optional AI uses an installed local Ollama model.');self.status.setWordWrap(True);box.addWidget(self.status)
        self.controls()

    def error(self,message):QMessageBox.warning(self,'Document revision',str(message))

    def task(self,operation,callback):
        def finished(value,error):
            if error:self.error(error)
            else:
                try:callback(value)
                except Exception as exc:self.error(exc)
        if self.run_task:self.run_task(operation,finished)
        else:
            try:value=operation()
            except Exception as exc:finished(None,exc)
            else:finished(value,None)

    def new(self):
        dialog=NewRevisionDialog(self.library,self)
        if dialog.exec()!=QDialog.DialogCode.Accepted:return
        if not self.close_session(release=False):return
        old,template,folder,rev=dialog.old.text(),dialog.template.text(),dialog.folder.text(),dialog.revision.text()
        self.task(lambda:RevisionWorkspace.create(folder,old,template,rev),self.bind)

    def open(self):
        path,_=QFileDialog.getOpenFileName(self,'Open revision workspace','','Revision workspace (revision.json)')
        if path and self.close_session(release=False):self.task(lambda:RevisionWorkspace.open(Path(path).parent),self.bind)

    def bind(self,work):
        same=self.work is not None and self.work.root==work.root and self.workspace_lock is not None
        candidate=self.workspace_lock if same else QLockFile(str(work.root/'revision.lock'))
        if not same and not candidate.tryLock(0):raise ValueError('This revision workspace is open in another instance. Close it there before opening it here.')
        if not same:
            if self.workspace_lock:self.workspace_lock.unlock()
            self.workspace_lock=candidate
        self.work=work;self.location.setText(str(work.root));self.title.setText(work.state['title']);self.revision.setText(work.state['revision']);self.summary.setText(work.state['change_summary'])
        self.destination.clear()
        for dest in work.destinations:self.destination.addItem(dest['label'],dest['id'])
        self.destination.setCurrentIndex(self.destination.count()-1)
        source=work.resolve(work.state['source_file'])
        if source.suffix=='.pdf':
            from PySide6.QtPdf import QPdfDocument
            from PySide6.QtPdfWidgets import QPdfView
            self.pdf_buffer=QBuffer(self);self.pdf_buffer.setData(source.read_bytes());self.pdf_buffer.open(QIODevice.OpenModeFlag.ReadOnly);self.pdf_document=QPdfDocument(self);self.pdf_document.load(self.pdf_buffer)
            view=QPdfView();view.setDocument(self.pdf_document);view.setPageMode(QPdfView.PageMode.MultiPage);view.setZoomMode(QPdfView.ZoomMode.FitToWidth);old=self.source_tabs.widget(1);self.source_tabs.removeTab(1);old.deleteLater();self.source_tabs.addTab(view,'Original PDF pages')
        else:
            old=self.source_tabs.widget(1);self.source_tabs.removeTab(1);old.deleteLater();self.original=QTextBrowser();self.original.setHtml(document_html(self.original,source));self.source_tabs.addTab(self.original,'Whole document')
        self.refresh();self.refresh_preview()

    def original_pages(self):
        if not self.work:return
        source=self.work.resolve(self.work.state['source_file'])
        if source.suffix=='.pdf':self.source_tabs.setCurrentIndex(1);return
        work=self.work
        def convert():
            from ..word_export import word_pdf
            output=work.root/'original-preview.pdf';word_pdf(source,output);return output
        def show(path):
            from PySide6.QtPdf import QPdfDocument
            from PySide6.QtPdfWidgets import QPdfView
            if self.pdf_document:self.pdf_document.close();self.pdf_document.deleteLater()
            if self.pdf_buffer:self.pdf_buffer.close();self.pdf_buffer.deleteLater()
            self.pdf_buffer=QBuffer(self);self.pdf_buffer.setData(path.read_bytes());self.pdf_buffer.open(QIODevice.OpenModeFlag.ReadOnly);self.pdf_document=QPdfDocument(self);self.pdf_document.load(self.pdf_buffer)
            view=QPdfView();view.setDocument(self.pdf_document);view.setPageMode(QPdfView.PageMode.MultiPage);view.setZoomMode(QPdfView.ZoomMode.FitToWidth);old=self.source_tabs.widget(1);self.source_tabs.removeTab(1);old.deleteLater();self.source_tabs.addTab(view,'Original pages');self.source_tabs.setCurrentIndex(1)
        self.task(convert,show)

    def refresh(self):
        if not self.work:return
        selected=self.sources.currentRow();placement=self.placements.currentRow();self.loading=True;self.sources.clear();self.placements.clear()
        for e in self.work.elements:
            label=f"[{self.work.status(e['id'])}] {e['kind'].title()}: "+(e['text'].replace('\n',' ')[:100] or 'Artwork')
            if e.get('warnings'):label='⚠ '+label
            self.sources.addItem(label);self.sources.item(self.sources.count()-1).setToolTip('\n'.join(e.get('warnings',[])))
        for p in self.work.state['placements']:
            dest=next(d['label'] for d in self.work.destinations if d['id']==p['destination']);self.placements.addItem(dest+' ← '+self.work.element(p['source'])['text'][:65])
        self.loading=False;self.sources.setCurrentRow(max(0,min(selected,len(self.work.elements)-1)));self.placements.setCurrentRow(placement)
        counts={status:sum(self.work.status(e['id'])==status for e in self.work.elements) for status in ('Not placed','Suggested','Approved','Excluded')}
        warnings=self.work.state['warnings'];self.status.setText(' · '.join(f'{n} {name.lower()}' for name,n in counts.items())+(f' · {len(warnings)} document warnings (hover to read)' if warnings else ''));self.status.setToolTip('\n'.join(warnings))
        self.controls();self.select_source(self.sources.currentRow())

    def controls(self):
        ready=bool(self.work);editable=ready and not self.work.state.get('authored');element=self.selected() if ready else None
        self.approve_button.setEnabled(bool(editable and element and not element.get('blocked')));self.exclude_button.setEnabled(bool(ready and element and (editable or self.work.status(element['id'])!='Approved')));self.ask_button.setEnabled(bool(editable and element));self.details_button.setEnabled(editable)
        self.text.setReadOnly(not editable);self.table.setEditTriggers(QTableWidget.EditTrigger.AllEditTriggers if editable else QTableWidget.EditTrigger.NoEditTriggers)

    def selected(self):
        index=self.sources.currentRow()
        return self.work.elements[index] if self.work and 0<=index<len(self.work.elements) else None

    def select_source(self,index):
        if self.loading or not self.work:return
        e=self.selected()
        if not e:return
        placement=next((p for p in self.work.state['placements'] if p['source']==e['id']),None);proposal=self.work.state['suggestions'].get(e['id'],{})
        content='<table border="1" cellpadding="5">'+''.join('<tr>'+''.join('<td>'+escape(cell)+'</td>' for cell in row)+'</tr>' for row in e['rows'])+'</table>' if e.get('rows') else '<p>'+escape(e['text']).replace('\n','<br>')+'</p>'
        self.source_view.setHtml('<h3>'+escape(e['kind'].title())+'</h3>'+content+''.join('<p><img width="350" src="'+QUrl.fromLocalFile(str(self.work.resolve(p))).toString()+'"></p>' for p in e.get('images',[]))+'<p>'+escape('\n'.join(e.get('warnings',[])))+'</p>')
        data=placement or proposal
        if data.get('destination'):self.destination.setCurrentIndex(self.destination.findData(data['destination']))
        self.text.setPlainText(data.get('text',e['text']) if e['kind']!='page' else data.get('text','') if placement else '')
        self.text.setEnabled(e['kind'] in {'text','heading','page'});rows=(placement or {}).get('rows') or e.get('rows') or []
        self.table.clear();self.table.setRowCount(len(rows));self.table.setColumnCount(max([len(r) for r in rows]+[0]))
        for ri,row in enumerate(rows):
            for ci,value in enumerate(row):self.table.setItem(ri,ci,QTableWidgetItem(value))
        self.table.resizeColumnsToContents();self.content.setCurrentIndex(1 if e['kind']=='table' else 0)
        self.editor_label.setText('Optional manual transcription: entering text transfers text instead of this page image.' if e['kind']=='page' else 'Edit text or table cells before approving. Artwork is retained as artwork.')
        self.proposal.setText(proposal.get('reason','Select a destination, then approve this element.')+(' — direct Word edits are active; return to placements to change transfers.' if self.work.state.get('authored') else ''));self.controls()

    def select_placement(self,index):
        if self.loading or not self.work or not 0<=index<len(self.work.state['placements']):return
        identity=self.work.state['placements'][index]['source'];self.sources.setCurrentRow(next(i for i,e in enumerate(self.work.elements) if e['id']==identity))

    def approve(self):
        e=self.selected()
        if not e:return
        try:
            rows=[[self.table.item(ri,ci).text() if self.table.item(ri,ci) else '' for ci in range(self.table.columnCount())] for ri in range(self.table.rowCount())] if e['kind']=='table' else None
            text=self.text.toPlainText() if e['kind'] in {'text','heading','page'} else e['text']
            self.work.approve(e['id'],self.destination.currentData(),text=text,rows=rows);self.refresh();self.refresh_preview()
        except Exception as exc:self.error(exc)

    def exclude(self):
        e=self.selected()
        if not e:return
        reason,ok=QInputDialog.getText(self,'Exclude source element','Reason for excluding it:')
        if ok:
            try:self.work.exclude(e['id'],reason);self.refresh();self.refresh_preview()
            except Exception as exc:self.error(exc)

    def unplace(self):
        if self.selected():
            try:self.work.unplace(self.selected()['id']);self.refresh();self.refresh_preview()
            except Exception as exc:self.error(exc)

    def move(self,offset):
        index=self.placements.currentRow()
        if self.work and 0<=index<len(self.work.state['placements']):
            try:self.work.move(self.work.state['placements'][index]['id'],offset);self.refresh();self.placements.setCurrentRow(max(0,index+offset));self.refresh_preview()
            except Exception as exc:self.error(exc)

    def undo(self):
        if self.work:self.work.undo();self.refresh();self.refresh_preview()

    def redo(self):
        if self.work:self.work.redo();self.refresh();self.refresh_preview()

    def save_details(self):
        if self.work:
            try:self.work.update_details(self.title.text(),self.revision.text(),self.summary.text());self.refresh_preview()
            except Exception as exc:self.error(exc)

    def refresh_preview(self):
        if not self.work:return
        try:
            path=build_document(self.work,self.work.root/'preview.docx');self.new_view.setHtml(document_html(self.new_view,path))
        except Exception as exc:self.new_view.setPlainText('Preview needs attention: '+str(exc))

    def ask(self):
        e=self.selected()
        if not e or self.work.state.get('authored'):return
        identity=e['id'];work=self.work;model=self.model.currentData();instruction=self.instruction.text()
        def apply(proposal):
            if self.work is work:work.suggest(identity,proposal['destination'],proposal['reason'],proposal.get('text'));self.refresh()
        self.task(lambda:LocalAssistant().propose(model,e,work.destinations,instruction) if model else suggest_placement(e,work.destinations,instruction),apply)

    def find_models(self):
        def apply(models):
            self.model.clear();self.model.addItem('Local structural matching',None)
            for model in models:self.model.addItem(model,model)
            self.status.setText('Choose an installed local model. Suggestions still require approval.' if models else 'No installed local models found. Structural matching is available.')
        self.task(lambda:LocalAssistant().models(),apply)

    def persist_document(self,data,source_key):
        atomic_bytes(self.work.root/'authored.docx',data);self.work.state['authored']=True;self.work.save();self.controls();return 'authored.docx'

    def open_word(self):
        if not self.work:return
        unresolved=sum(self.work.status(e['id']) in {'Not placed','Suggested'} for e in self.work.elements)
        if unresolved and QMessageBox.question(self,'Unresolved source elements',f'{unresolved} source elements are not approved or excluded. Open the current draft anyway?',QMessageBox.StandardButton.Yes|QMessageBox.StandardButton.No,QMessageBox.StandardButton.No)!=QMessageBox.StandardButton.Yes:return
        try:
            temporary=tempfile.TemporaryDirectory(prefix='sop-revision-word-');path=build_document(self.work,Path(temporary.name)/'revision.docx')
            self.word_editor.open_document(path,temporary,None,'');self.right_tabs.setCurrentIndex(3)
        except Exception as exc:self.error(exc)

    def return_to_placements(self):
        if not self.work:return
        if QMessageBox.question(self,'Return to placements','Keep the directly edited document as a backup and rebuild from approved placements? Manual Word changes will not be in the rebuilt draft.',QMessageBox.StandardButton.Yes|QMessageBox.StandardButton.No,QMessageBox.StandardButton.No)!=QMessageBox.StandardButton.Yes:return
        try:
            self.word_editor.close_document()
            if (self.work.root/'authored.docx').exists():
                import uuid
                atomic_bytes(self.work.root/('manual-backup-'+uuid.uuid4().hex[:8]+'.docx'),(self.work.root/'authored.docx').read_bytes())
            self.work.state['authored']=False;self.work.save();self.refresh();self.refresh_preview();self.right_tabs.setCurrentIndex(0)
        except Exception as exc:self.error(exc)

    def export(self):
        if not self.work:return
        try:
            self.word_editor.save_document();unresolved=sum(self.work.status(e['id']) in {'Not placed','Suggested'} for e in self.work.elements)
            if unresolved and QMessageBox.question(self,'Incomplete transfer',f'{unresolved} source elements still need a decision. Export the approved draft with its report?',QMessageBox.StandardButton.Yes|QMessageBox.StandardButton.No,QMessageBox.StandardButton.No)!=QMessageBox.StandardButton.Yes:return
            path,_=QFileDialog.getSaveFileName(self,'Save revised Word document',str(self.work.root/f'Rev {self.work.state["revision"]}.docx'),'Word document (*.docx)')
            if path:build_document(self.work,path);self.status.setText('Saved revised Word document and transfer report: '+path)
        except Exception as exc:self.error(exc)

    def close_session(self,release=True):
        try:
            self.word_editor.close_document()
            if self.work:self.work.save()
            if self.pdf_document:self.pdf_document.close();self.pdf_document.deleteLater();self.pdf_document=None
            if self.pdf_buffer:self.pdf_buffer.close();self.pdf_buffer.deleteLater();self.pdf_buffer=None
            if release and self.workspace_lock:self.workspace_lock.unlock();self.workspace_lock=None
            return True
        except Exception as exc:self.error(exc);return False
