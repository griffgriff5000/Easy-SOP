import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from PIL import Image
from pypdf import PdfReader
from docx import Document
import pytest
from PySide6.QtWidgets import QApplication

from sop_capture.models import Annotation, Project, Step
from sop_capture.storage import ProjectStore
from sop_capture.imaging import render_step
from sop_capture.pdf_export import export_pdf
from sop_capture.word_export import export_docx
from sop_capture.ui.main_window import MainWindow
from sop_capture.review import review_project


def test_callout_pixels_persistence_and_redaction(tmp_path):
    store = ProjectStore(tmp_path)
    step = Step(title="Open", section="Preparation", annotations=[Annotation(kind="callout", points=[40, 140, 140, 30], color="#0000ff", text="Open the menu", font_size=18)])
    step.image = store.add_image(Image.new("RGB", (500, 250), "#eeeeee"), step.id)
    project = Project(steps=[step])
    store.save(project)
    reopened = store.load()
    assert reopened.schema_version == 3
    assert reopened.steps[0].section == "Preparation"
    image = render_step(tmp_path, reopened.steps[0], 1)
    assert image.getpixel((40, 140)) == (0, 0, 255)
    assert image.getpixel((150, 40)) != (238, 238, 238)
    step.annotations.append(Annotation(kind="redact", points=[20, 120, 60, 160]))
    assert render_step(tmp_path, step, 1).getpixel((40, 140)) == (0, 0, 0)


def test_callout_label_and_target_edit_independently(tmp_path):
    app = QApplication.instance() or QApplication([])
    store = ProjectStore(tmp_path)
    step = Step(annotations=[Annotation(kind="callout", points=[40, 180, 200, 40], text="Choose this", font_size=20)])
    step.image = store.add_image(Image.new("RGB", (600, 300), "white"), step.id)
    window = MainWindow()
    window.set_project(Project(steps=[step]), store)
    window.canvas.set_tool("edit")
    window.canvas.apply_gesture((240, 55), (270, 75))
    assert window.current_step().annotations[0].points == [40, 180, 230, 60]
    window.canvas.apply_gesture((40, 180), (60, 200))
    assert window.current_step().annotations[0].points == [60, 200, 230, 60]
    window.undo()
    assert window.current_step().annotations[0].points == [40, 180, 230, 60]
    window.close()


def test_group_range_rename_undo_and_export_headings(tmp_path):
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    store = ProjectStore(tmp_path)
    window.set_project(Project(steps=[Step(title="One"), Step(title="Two"), Step(title="Three")]), store)
    window.group_steps("Preparation", 0, 1)
    window.group_steps("Submission", 2, 2)
    window.rename_section("Preparation", "Before starting")
    assert [s.section for s in window.project.steps] == ["Before starting", "Before starting", "Submission"]
    window.undo()
    assert window.project.steps[0].section == "Preparation"
    window.redo()
    window.flush_save()
    assert store.load().steps[2].section == "Submission"
    pdf = PdfReader(export_pdf(window.project, tmp_path, tmp_path / "groups.pdf"))
    text = "".join(page.extract_text() for page in pdf.pages)
    assert text.count("Before starting") == 1 and text.count("Submission") == 1
    doc = Document(export_docx(window.project, tmp_path, tmp_path / "groups.docx"))
    assert [p.text for p in doc.paragraphs].count("Before starting") == 1
    assert "Submission" in [p.text for p in doc.paragraphs]
    window.close()


def test_review_reports_missing_fields_and_step_issues_without_mutating(tmp_path):
    project = Project(title="", author="", purpose="", scope="", steps=[Step(title="New step", description="", image="Screenshots/missing.png")])
    before = project.to_dict()
    issues = review_project(project, tmp_path)
    assert any(issue.field == "title" for issue in issues)
    assert any(issue.field == "description" and issue.step_id == project.steps[0].id for issue in issues)
    assert any(issue.field == "image" for issue in issues)
    assert project.to_dict() == before


def test_complete_review_respects_company_default_and_valid_images(tmp_path):
    store = ProjectStore(tmp_path)
    step = Step(title="Open dashboard", description="Select Dashboard in the menu.")
    step.image = store.add_image(Image.new("RGB", (100, 100), "white"), step.id)
    project = Project(title="Permit process", author="Author", purpose="Prepare a permit", scope="Site permits", steps=[step],
                      template={"name":"Company", "kind":"custom", "settings":{"company":"Company name"}})
    assert review_project(project, tmp_path) == []


def test_old_projects_have_empty_section():
    raw = Project(steps=[Step()]).to_dict()
    raw["steps"][0].pop("section")
    assert Project.from_dict(raw).steps[0].section == ""


def test_review_can_return_to_step_and_cancel_export(tmp_path, monkeypatch):
    from PySide6.QtCore import QTimer
    from PySide6.QtWidgets import QListWidget, QPushButton, QFileDialog
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    window.set_project(Project(steps=[Step(title="Choose Dashboard")]), ProjectStore(tmp_path))
    def fix_description():
        dialog = app.activeModalWidget()
        items = dialog.findChild(QListWidget)
        row = next(i for i in range(items.count()) if 'add instructions' in items.item(i).text())
        items.setCurrentRow(row)
        next(b for b in dialog.findChildren(QPushButton) if b.text() == 'Fix selected item').click()
    QTimer.singleShot(0, fix_description)
    assert not window.show_review(for_export=True)
    assert window.steps_list.currentRow() == 0
    monkeypatch.setattr(window, 'show_review', lambda **kwargs: False)
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args: pytest.fail('Cancelled review must not open file picker'))
    window.export_dialog()
    window.export_word_dialog()
    window.close()


def test_review_flags_missing_assets(tmp_path):
    project = Project(logo='Assets/missing.png')
    assert any(issue.field == 'logo' for issue in review_project(project, tmp_path))


def test_sections_stay_with_steps_in_page_per_step_exports(tmp_path):
    project = Project(steps=[Step(title='Open dashboard', section='Preparation', description='Open the menu.')],
                      template={'name': 'Paged', 'kind': 'custom', 'settings': {'page_per_step': True}})
    pdf = PdfReader(export_pdf(project, tmp_path, tmp_path / 'paged.pdf'))
    texts = [page.extract_text() for page in pdf.pages]
    assert any('Preparation' in text and 'Open dashboard' in text for text in texts)
    document = Document(export_docx(project, tmp_path, tmp_path / 'paged.docx'))
    section = next(p for p in document.paragraphs if p.text == 'Preparation')
    title = next(p for p in document.paragraphs if 'Open dashboard' in p.text)
    assert section.paragraph_format.page_break_before
    assert not title.paragraph_format.page_break_before


def test_callout_at_edge_moves_from_visible_label_position(tmp_path):
    from sop_capture.imaging import callout_layout
    app = QApplication.instance() or QApplication([])
    store = ProjectStore(tmp_path)
    annotation = Annotation(kind='callout', points=[40, 180, 590, 40], text='Choose this', font_size=20)
    step = Step(annotations=[annotation])
    step.image = store.add_image(Image.new('RGB', (600, 300), 'white'), step.id)
    window = MainWindow()
    window.set_project(Project(steps=[step]), store)
    box, _, _ = callout_layout(annotation, (600, 300))
    moved = window.canvas.edit_points(annotation, None, (box[0]+20, 55), (box[0], 55))
    assert moved[2] == box[0]-20
    window.close()
