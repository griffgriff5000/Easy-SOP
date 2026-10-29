import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from copy import deepcopy
from pathlib import Path
import pytest
from docx import Document
from PySide6.QtWidgets import QApplication
from sop_capture.models import Project,Step
from sop_capture.storage import ProjectStore
from sop_capture.templates import TemplateLibrary
from sop_capture.word_export import export_docx
from sop_capture.template_text import template_paragraphs,edit_snapshot

@pytest.fixture(scope='module',autouse=True)
def app():return QApplication.instance() or QApplication([])

def source_document(path):
    document=Document();document.sections[0].header.paragraphs[0].text='Controlled copy'
    p=document.add_paragraph();p.add_run('Owner: ');p.add_run('Old owner').bold=True;p.add_run(' - Operations').italic=True
    table=document.add_table(rows=1,cols=2);table.cell(0,0).text='Reference';table.cell(0,1).text='Old reference'
    document.add_paragraph('{{title}}');document.add_paragraph('{{steps}}');document.save(path)
    return path

def test_template_text_edits_preserve_structure_runs_and_original(tmp_path):
    source=source_document(tmp_path/'company.docx');original=source.read_bytes()
    records=template_paragraphs(source)
    owner=next(r for r in records if r['text'].startswith('Owner:'))
    reference=next(r for r in records if r['text']=='Old reference')
    patches=edit_snapshot(source,{owner['key']:'Owner: William - Operations',reference['key']:'REF-123'})
    library=TemplateLibrary(tmp_path/'library');profile=library.save('Company',{},source_docx=source);root=tmp_path/'project'
    project=Project(title='My SOP',steps=[Step(title='Open dashboard')],template=library.snapshot(profile['id'],root),template_text_edits=patches)
    store=ProjectStore(root);store.save(project);assert store.load().schema_version==6
    document=Document(export_docx(project,root,tmp_path/'edited.docx'))
    assert source.read_bytes()==original
    assert document.tables[0].cell(0,1).text=='REF-123'
    paragraph=next(p for p in document.paragraphs if p.text.startswith('Owner:'))
    assert paragraph.text=='Owner: William - Operations'
    assert any(r.bold and 'William' in r.text for r in paragraph.runs)
    assert any(r.italic and 'Operations' in r.text for r in paragraph.runs)
    assert document.sections[0].header.paragraphs[0].text=='Controlled copy'
    assert any('Open dashboard' in p.text for p in document.paragraphs)

def test_template_patch_rejects_wrong_source_and_dynamic_fields(tmp_path):
    source=source_document(tmp_path/'source.docx');records=template_paragraphs(source)
    dynamic=next(r for r in records if r['text']=='{{title}}')
    assert not dynamic['editable']
    with pytest.raises(ValueError):edit_snapshot(source,{dynamic['key']:'Replaced title'})
    key=next(r['key'] for r in records if r['text']=='Old reference');patches=edit_snapshot(source,{key:'New'})
    library=TemplateLibrary(tmp_path/'library');profile=library.save('Company',{},source_docx=source);root=tmp_path/'project'
    project=Project(template=library.snapshot(profile['id'],root),template_text_edits=patches)
    project.template_text_edits['source_hash']='0'*64
    with pytest.raises(ValueError,match='changed'):export_docx(project,root,tmp_path/'bad.docx')

def test_inline_preview_uses_snapshot_and_drops_old_project_result(tmp_path):
    import time,threading
    from reportlab.pdfgen import canvas
    from sop_capture.ui.builder_preview import BuilderPreview,document_key
    app=QApplication.instance();first=Project(title='First');second=Project(title='Second');current=[first]
    gate=threading.Event()
    def export(project,root,path):
        gate.wait(3);pdf=canvas.Canvas(str(path));pdf.drawString(40,700,project.title);pdf.save();return path
    preview=BuilderPreview(lambda:document_key(current[0],tmp_path),exporter=export)
    preview.request(first,tmp_path);assert preview.running
    preview.reset();current[0]=second;preview.request(second,tmp_path);gate.set()
    deadline=time.monotonic()+5
    while preview.running and time.monotonic()<deadline:app.processEvents();time.sleep(.01)
    assert not preview.running
    assert preview.loaded_key==document_key(second,tmp_path)
    assert preview.document.pageCount()==1
    preview.clear_document();preview.close()


def test_builder_shows_selected_template_and_preserves_edits_with_undo(tmp_path):
    from sop_capture.ui.main_window import MainWindow
    source=source_document(tmp_path/'company.docx');library=TemplateLibrary(tmp_path/'library');profile=library.save('Company layout',{},source_docx=source)
    window=MainWindow(template_library=library);store=ProjectStore(tmp_path/'project');window.set_project(Project(steps=[Step(title='First step')]),store)
    window.apply_template(profile['id'])
    assert 'Company layout' in window.builder_template_label.text()
    editor=window.template_text_editor
    index=next(i for i,row in enumerate(editor.records) if row['text']=='Old reference');editor.picker.setCurrentIndex(index);editor.text.setPlainText('REF-456')
    assert window.project.template_text_edits
    window.flush_save();assert store.load().template_text_edits==window.project.template_text_edits
    window.undo();assert not window.project.template_text_edits
    window.redo();assert window.project.template_text_edits
    window.apply_template('standard');assert not window.project.template_text_edits
    window.undo();assert window.project.template_text_edits
    window.close()

def test_complex_word_reference_is_protected(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    path=tmp_path/'references.docx';document=Document();paragraph=document.add_paragraph();run=paragraph.add_run('Company note')
    note=OxmlElement('w:footnoteReference');note.set(qn('w:id'),'1');run._r.append(note);document.add_paragraph('{{steps}}');document.save(path)
    record=next(row for row in template_paragraphs(path) if row['text']=='Company note')
    assert not record['editable']
    with pytest.raises(ValueError):edit_snapshot(path,{record['key']:'Updated company note'})


def test_ordinary_edits_do_not_install_stale_render(tmp_path):
    import threading,time
    from reportlab.pdfgen import canvas
    from sop_capture.ui.builder_preview import BuilderPreview,document_key
    app=QApplication.instance();project=Project(title='Before');gate=threading.Event()
    def export(project,root,path):
        gate.wait(3);pdf=canvas.Canvas(str(path));pdf.drawString(40,700,project.title);pdf.save();return path
    preview=BuilderPreview(lambda:document_key(project,tmp_path),exporter=export)
    preview.request(project,tmp_path);project.title='After';preview.invalidate();gate.set()
    deadline=time.monotonic()+5
    while preview.running and time.monotonic()<deadline:app.processEvents();time.sleep(.01)
    assert preview.document is None
    assert 'changed' in preview.status.text().lower()
    preview.close()

def test_builder_fields_update_metadata_without_leaving_preview(tmp_path):
    from sop_capture.ui.main_window import MainWindow
    window=MainWindow();window.set_project(Project(title='My SOP',reference='OLD',steps=[Step(title='Action')]),ProjectStore(tmp_path))
    window.builder_metadata_fields['reference'].setText('NEW-REF')
    assert window.project.reference=='NEW-REF'
    assert window.document_fields['reference'].text()=='NEW-REF'
    window.document_fields['author'].setText('Author')
    assert window.builder_metadata_fields['author'].text()=='Author'
    window._reveal_builder_field('reference')
    assert window.document_tabs.currentIndex()==0 and window.builder_edit_tabs.currentIndex()==0
    window.close()


def test_template_cannot_be_overwritten_or_add_sop_markers(tmp_path):
    source=source_document(tmp_path/'company.docx');records=template_paragraphs(source)
    key=next(row['key'] for row in records if row['text']=='Old reference')
    with pytest.raises(ValueError,match='markers'):edit_snapshot(source,{key:'{{steps}}'})
    library=TemplateLibrary(tmp_path/'library');profile=library.save('Company',{},source_docx=source);root=tmp_path/'project';project=Project(template=library.snapshot(profile['id'],root))
    path=ProjectStore(root).resolve(project.template['docx']);before=path.read_bytes()
    with pytest.raises(ValueError,match='separate export'):export_docx(project,root,path)
    assert path.read_bytes()==before

def test_closing_stops_capture_before_waiting_for_preview():
    from sop_capture.ui.main_window import MainWindow
    from PySide6.QtGui import QCloseEvent
    window=MainWindow()
    class Session:state='recording'
    window.session=Session();window.builder_preview.running=True;called=[]
    def stop():called.append(True);window.session=None
    window.stop_capture=stop;event=QCloseEvent();window.closeEvent(event)
    assert called and not event.isAccepted()
    window.builder_preview.running=False;window.close()
