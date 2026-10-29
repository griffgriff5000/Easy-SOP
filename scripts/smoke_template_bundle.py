"""Check real document generation inside a frozen application subprocess."""
from pathlib import Path
import multiprocessing as mp
import multiprocessing.spawn
import sys
import tempfile
import textwrap

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from sop_capture.models import Project, Step, Annotation
from sop_capture.designs import new_design, new_layer, import_artwork
from sop_capture.storage import ProjectStore
from sop_capture.pdf_export import export_pdf as native_pdf
from PIL import Image
from sop_capture.templates import TemplateLibrary
from sop_capture.word_export import create_example, export_docx
from sop_capture.document_export import export_pdf


def main():
    executable = str(Path(sys.argv[1]).resolve())
    mp.set_executable(executable)
    mp.spawn.get_command_line = lambda **kw: [executable, "--multiprocessing-fork", *[f"{key}={value}" for key, value in kw.items()]]
    preparation = mp.spawn.get_preparation_data
    def frozen_preparation(name):
        data = preparation(name)
        # Keep the frozen child's own runtime-hook paths (including pywin32),
        # rather than replacing them with the source parent's paths.
        data.pop("sys_path", None)
        data.pop("init_main_from_name", None)
        data.pop("init_main_from_path", None)
        return data
    mp.spawn.get_preparation_data = frozen_preparation
    with tempfile.TemporaryDirectory(prefix="sop-frozen-template-") as temporary:
        root = Path(temporary)
        example = create_example(root / "source.docx")
        library = TemplateLibrary(root / "library")
        profile = library.save("Bundle check", {}, source_docx=example)
        project = Project(title="Frozen template export", steps=[Step(title="Check output", description="Synthetic package verification.")], template=library.snapshot(profile["id"], root / "project"))
        project.company_information = "<h2>Company requirements</h2><p><b>Owner:</b> Example team. <i>Approval required.</i></p><table><tr><td>Reference</td><td>RA-001</td></tr></table>"
        from sop_capture.template_text import template_paragraphs, edit_snapshot
        editable = next(row for row in template_paragraphs(example) if row["editable"] and row["text"])
        project.template_text_edits = edit_snapshot(example, {editable["key"]: editable["text"]+" (edited in Builder)"})
        # Exercise custom vectors, text/fonts and embedded artwork in the frozen renderer.
        artwork = root / "art.png"
        Image.new("RGBA", (60, 40), (10, 120, 220, 160)).save(artwork)
        design = new_design(400, 240)
        design["layers"] = [new_layer("rect", x=10, y=10, w=350, h=200, fill="#ffffff"),
                            new_layer("text", x=30, y=30, w=300, h=80, text="Custom annotation", font="vera", bold=True),
                            new_layer("arrow", x=20, y=160, w=200, h=1), import_artwork(artwork)]
        project.steps[0].image = ProjectStore(root / "project").add_image(Image.new("RGB", (600, 400), "white"), project.steps[0].id)
        project.steps[0].annotations = [Annotation(kind="design", points=[20,20,420,260], design=design), Annotation(kind="magnifier", points=[430,280,570,380], zoom=2.5)]
        ProjectStore(root / "project").save(project)
        for target, extension in ((export_docx, ".docx"), (export_pdf, ".pdf"), (native_pdf, ".native.pdf")):
            output = root / ("result"+extension)
            process = mp.get_context("spawn").Process(target=target, args=(project, root / "project", output))
            process.start()
            process.join(60)
            if process.is_alive():
                process.terminate()
                process.join()
                raise RuntimeError("Frozen document generation timed out")
            if process.exitcode != 0 or not output.exists(): raise RuntimeError("Frozen document export failed: "+extension)
            print("Packaged template export:", extension, "passed", output.stat().st_size, "bytes", flush=True)
        from sop_capture.authored_document import store_document, source_fingerprint
        from docx import Document
        edited=root/'edited.docx';document=Document(root/'result.docx');document.add_paragraph('Saved native editor content');document.save(edited)
        project.document_file=store_document(ProjectStore(root/'project'),edited.read_bytes())
        project.document_source_key=source_fingerprint(project)
        for target,extension in ((export_docx,'.docx'),(export_pdf,'.pdf')):
            output=root/('authored'+extension)
            process=mp.get_context('spawn').Process(target=target,args=(project,root/'project',output));process.start();process.join(60)
            if process.is_alive():process.terminate();process.join();raise RuntimeError('Frozen authored export timed out')
            if process.exitcode!=0 or not output.exists():raise RuntimeError('Frozen authored export failed')
            if extension=='.docx':assert output.read_bytes()==edited.read_bytes()
            else:
                from pypdf import PdfReader
                assert 'Saved native editor content' in ''.join(p.extract_text() for p in PdfReader(output).pages)
            print('Packaged edited-document export:',extension,'passed',flush=True)
        # Builtins.exec is available in a frozen spawn child without importing
        # this external verification script into the packaged application.
        code = f'''
from pathlib import Path
import time
from PySide6.QtWidgets import QApplication
from sop_capture.ui.word_editor import WordEditor
from sop_capture.authored_document import store_document
from sop_capture.storage import ProjectStore
from docx import Document
app=QApplication([])
root=Path({str(root)!r})
store=ProjectStore(root/'project')
project=store.load()
def persist(data,key):
    project.document_file=store_document(store,data)
    store.save(project)
    return project.document_file
editor=WordEditor(persist);editor.resize(1100,750);editor.show()
try:
    editor.open_document(root/'edited.docx',None,None,'')
    assert editor.container is not None
    deadline=time.monotonic()+1
    while time.monotonic()<deadline:app.processEvents();time.sleep(.02)
    selection=editor.session.document.ActiveWindow.Selection
    selection.EndKey(6);selection.TypeText('\\rFrozen native Word editing works.')
    editor.save_document()
    assert any('Frozen native Word editing works' in p.text for p in Document(store.resolve(project.document_file)).paragraphs)
    (root/'native-editor-passed.txt').write_text('passed')
finally:
    editor.close_document(save=False);editor.close();app.processEvents()
from sop_capture.ui.templates_page import TemplatesPage
from sop_capture.templates import TemplateLibrary
from sop_capture.models import Project,Step
from sop_capture.word_export import export_docx
library=TemplateLibrary(root/'prepared-library')
page=TemplatesPage(library,lambda identity:None,lambda:None)
profile=page.import_path(str(root/'source.docx'),'Prepared company template')
assert profile['prepared']
prepared=Project(title='Prepared template SOP',steps=[Step(title='Generated clean step')],template=library.snapshot(profile['id'],root/'prepared-project'))
export_docx(prepared,root/'prepared-project',root/'prepared-result.docx')
text='\\n'.join(p.text for p in Document(root/'prepared-result.docx').paragraphs)
assert 'Generated clean step' in text and 'Prepared template SOP' in text
page.close();page.deleteLater();app.processEvents()
(root/'prepared-template-passed.txt').write_text('passed')
from sop_capture.revisions import RevisionWorkspace
from sop_capture.revision_import import inspect_source
from sop_capture.revision_export import build_document
from sop_capture.ui.revision_page import RevisionPage
pdf_elements,pdf_warnings=inspect_source(root/'result.pdf',root/'pdf-revision-assets')
assert any(e['kind']=='text' for e in pdf_elements)
assert any(e['kind']=='page' for e in pdf_elements)
revision=RevisionWorkspace.create(root/'revision-workspace',root/'edited.docx',root/'source.docx','B')
element=next(e for e in revision.elements if e['kind']=='text' and e['text'])
revision.approve(element['id'],'end',text='Approved revision transfer')
from sop_capture.ui.main_window import MainWindow
from sop_capture.workspace_layouts import LayoutLibrary
from sop_capture.recent_projects import RecentProjects
from sop_capture.designs import AnnotationLibrary
window=MainWindow(TemplateLibrary(root/'revision-library'),RecentProjects(root/'revision-recent.json'),AnnotationLibrary(root/'revision-annotations'),LayoutLibrary(root/'revision-layouts.json'))
review=window.revision_page;review.bind(revision);window.show_page(7);window.resize(1300,900);window.show();review.right_tabs.setCurrentIndex(3);app.processEvents()
draft=build_document(revision,root/'revision-draft.docx')
review.word_editor.open_document(draft,None,None,'')
selection=review.word_editor.session.document.ActiveWindow.Selection
selection.EndKey(6);selection.TypeText('\\rRevision native editing preserved.')
review.word_editor.save_document();assert revision.state['authored']
window.workspaces.hosts['revision'].panels['properties'].setFloating(True);app.processEvents()
window.show_page(5);app.processEvents();assert not window.workspaces.hosts['revision'].panels['properties'].isVisible()
window.show_page(7);app.processEvents();review.word_editor.save_document()
assert window.close();app.processEvents()
result=build_document(revision,root/'revision-B.docx')
assert any('Revision native editing preserved.' in p.text for p in Document(result).paragraphs)
(root/'revision-passed.txt').write_text('passed')
'''
        code="import faulthandler\n_crash=open("+repr(str(root/'native-editor-crash.txt'))+",'w')\nfaulthandler.enable(_crash)\ntry:\n"+textwrap.indent(code,'    ')+"\nexcept BaseException:\n    import traceback\n    Path("+repr(str(root/'native-editor-failure.txt'))+").write_text(traceback.format_exc())\n    raise\n"
        process=mp.get_context('spawn').Process(target=exec,args=(code,{'__name__':'__native_smoke__'}));process.start();process.join(60)
        if process.is_alive():process.terminate();process.join();raise RuntimeError('Frozen Word editor timed out')
        if process.exitcode!=0 or not (root/'native-editor-passed.txt').exists():
            for name in ('native-editor-crash.txt','native-editor-failure.txt'):
                if (root/name).exists():print((root/name).read_text(),flush=True)
            raise RuntimeError(f'Frozen native Word editor failed, exit {process.exitcode}')
        print('Packaged native Word editor: embedded edit and project save passed',flush=True)
        if not (root/'prepared-template-passed.txt').exists():raise RuntimeError('Frozen template preparation failed')
        print('Packaged automatic Word template preparation: passed',flush=True)
        if not (root/'revision-passed.txt').exists():raise RuntimeError('Frozen revision transfer failed')
        print('Packaged revision: PDF import, approved transfer and native Word edits passed',flush=True)


if __name__ == "__main__": main()
