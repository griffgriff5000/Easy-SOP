from pathlib import Path
import zipfile
import pytest
from PIL import Image
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from sop_capture.models import Project, Step, Annotation
from sop_capture.storage import ProjectStore
from sop_capture.templates import TemplateLibrary
from sop_capture.word_export import inspect_docx, export_docx


def make_template(path):
    doc = Document()
    doc.styles.add_style("Company instruction", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("Company master", style="Title")
    title = doc.add_paragraph()
    title.add_run("{{ti")
    title.add_run("tle}}")
    doc.sections[0].header.paragraphs[0].text = "Controlled {{reference}}"
    doc.add_paragraph("{{steps}}")
    doc.save(path)


def test_word_styles_tokens_and_redacted_screenshots(tmp_path):
    source = tmp_path / "company.docx"
    make_template(source)
    info = inspect_docx(source)
    assert "Company instruction" in info["styles"]
    assert info["has_steps"]
    library = TemplateLibrary(tmp_path / "library")
    profile = library.save("Company Word", {"body_style": "Company instruction"}, source_docx=source)
    store = ProjectStore(tmp_path / "project")
    step = Step(title="Open dashboard", description="Select <permit> & continue", annotations=[Annotation(kind="redact", points=[10, 10, 40, 40])])
    step.image = store.add_image(Image.new("RGB", (100, 80), "white"), step.id)
    project = Project(title="Permit procedure", reference="PV-42", steps=[step], template=library.snapshot(profile["id"], store.root))
    target = export_docx(project, store.root, tmp_path / "result.docx")
    doc = Document(target)
    assert "Permit procedure" in [p.text for p in doc.paragraphs]
    assert "Controlled PV-42" in doc.sections[0].header.paragraphs[0].text
    instruction = next(p for p in doc.paragraphs if "Select <permit>" in p.text)
    assert instruction.style.name == "Company instruction"
    assert len(doc.inline_shapes) == 1
    with zipfile.ZipFile(target) as archive:
        from io import BytesIO
        image = Image.open(BytesIO(archive.read(next(n for n in archive.namelist() if n.startswith('word/media/')))))
        assert image.getpixel((20, 20)) == (0, 0, 0)
    assert "{{title}}" in Document(source).paragraphs[1].text


def test_missing_anchor_is_disclosed_and_content_appended(tmp_path):
    source = tmp_path / "plain.docx"
    doc = Document()
    doc.add_paragraph("Existing company content")
    doc.save(source)
    assert not inspect_docx(source)["has_steps"]
    store = ProjectStore(tmp_path / "project")
    library = TemplateLibrary(tmp_path / "library")
    profile = library.save("Plain", {}, source_docx=source)
    project = Project(steps=[Step(title="Appended action")], template=library.snapshot(profile["id"], store.root))
    result = Document(export_docx(project, store.root, tmp_path / "append.docx"))
    assert result.paragraphs[0].text == "Existing company content"
    assert any("Appended action" in p.text for p in result.paragraphs)


def test_invalid_archive_rejected(tmp_path):
    file = tmp_path / "bad.docx"
    file.write_bytes(b"not a Word file")
    with pytest.raises(ValueError): inspect_docx(file)


@pytest.mark.parametrize("pieces", [("INCLUDE", 'TEXT "C:\\private.txt"'), ("DDE", 'AUTO server topic')])
def test_split_active_word_fields_are_rejected(tmp_path, pieces):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    p = doc.add_paragraph()
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    p.add_run()._r.append(start)
    for piece in pieces:
        instruction = OxmlElement("w:instrText")
        instruction.text = piece
        p.add_run()._r.append(instruction)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    p.add_run()._r.append(end)
    path = tmp_path / "active.docx"
    doc.save(path)
    with pytest.raises(ValueError, match="fields"): inspect_docx(path)


def test_header_logo_relationship_belongs_to_header(tmp_path):
    from lxml import etree
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "{{logo}}"
    doc.add_paragraph("{{steps}}")
    source = tmp_path / "header.docx"
    doc.save(source)
    logo = tmp_path / "logo.png"
    Image.new("RGB", (100, 50), "blue").save(logo)
    library = TemplateLibrary(tmp_path / "library")
    profile = library.save("Header", {}, source_docx=source, logo=logo)
    project = Project(template=library.snapshot(profile["id"], tmp_path / "project"))
    output = export_docx(project, tmp_path / "project", tmp_path / "with-logo.docx")
    with zipfile.ZipFile(output) as archive:
        header = etree.fromstring(archive.read("word/header1.xml"))
        reference = header.xpath('//@r:embed', namespaces={'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'})[0]
        relationships = etree.fromstring(archive.read("word/_rels/header1.xml.rels"))
        assert any(r.get("Id") == reference and r.get("Type").endswith("/image") for r in relationships)
